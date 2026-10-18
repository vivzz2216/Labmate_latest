import importlib.util
import json
import tempfile
from pathlib import Path

from docx import Document
from PIL import Image
import pytest

spec=importlib.util.spec_from_file_location('manual_test_report',Path(__file__).parents[1]/'scripts/build_manual_test_report.py')
report=importlib.util.module_from_spec(spec)
spec.loader.exec_module(report)


def fixture(root,success=True):
    image=root/'test_1.png';Image.new('RGB',(500,300),'white').save(image)
    (root/'extracted.json').write_text(json.dumps([{'id':1,'language':'html','text':'Experiment 1: Registration\nProblem Statement\nValidate the form.'}]))
    (root/'run.json').write_text(json.dumps({'manual':'registration.docx','sha256':'test-source-hash'}))
    execution={'success':success,'error':'' if success else 'Wrong visible result','screenshots':[str(image),str(image)],'verification':{'status':'passed' if success else 'failed','coverage':{'validation':'passed' if success else 'unverified'},'tests':[{'name':'Invalid input','status':'passed' if success else 'failed','assertions':[{'status':'passed'}] if success else []}]}}
    (root/'experiment_01.json').write_text(json.dumps({'id':1,'execution':execution,'project':{'project_files':{'README.md':'Open index.html in a browser.'}}}))
    return image


def test_every_capture_occurrence_is_embedded_even_when_pixels_are_identical():
    with tempfile.TemporaryDirectory() as tmp:
        root=Path(tmp);fixture(root);path=root/'results.docx'
        report.build(root,path)
        doc=Document(path)
        assert len(doc.inline_shapes)==2
        text='\n'.join(p.text for p in doc.paragraphs)
        assert '1 of 1 experiments passed.' in text
        assert '2 genuine browser captures embedded' in text
        assert 'no successful api-generation claim' in text.lower()


def test_failed_execution_is_reported_as_failed():
    with tempfile.TemporaryDirectory() as tmp:
        root=Path(tmp);fixture(root,False);path=root/'results.docx'
        report.build(root,path)
        text='\n'.join(p.text for p in Document(path).paragraphs)
        assert '0 of 1 experiments passed.' in text
        assert 'Wrong visible result' in text


def test_missing_image_prevents_incomplete_export():
    with tempfile.TemporaryDirectory() as tmp:
        root=Path(tmp);fixture(root).unlink()
        with pytest.raises(ValueError,match='Missing required screenshot'):
            report.build(root,root/'incomplete.docx')
        assert not (root/'incomplete.docx').exists()


def test_workflow_word_export_rechecks_verified_evidence_at_write_time():
    from types import SimpleNamespace
    from app.services.assignment_workflow import build_word_report
    with tempfile.TemporaryDirectory() as tmp:
        root=Path(tmp);original=root/'manual.docx';Document().save(original)
        upload=SimpleNamespace(file_type='docx',file_path=str(original),original_filename='manual.docx')
        workflow=SimpleNamespace(id=199,output_name='verified',instructions='',results=[{
            'id':1,'question':'Show an HTML page','language':'html','status':'completed',
            'verification':{'status':'passed','tests':[]},'screenshot_paths':[str(root/'deleted.png')]}])
        with pytest.raises(ValueError,match='required screenshot is missing'):
            build_word_report(upload,workflow)
