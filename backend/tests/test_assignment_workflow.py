"""Regression checks for preserving user documents and reporting real results."""

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from PIL import Image

from app.services.assignment_workflow import build_word_report, document_text, validate_generated_code


class AssignmentReportTests(unittest.TestCase):
    def test_diagnostic_word_handles_ansi_logs_and_xml_controls_without_mutating_results(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);original=Document();original.add_paragraph('Original manual');source=root/'manual.docx';original.save(source)
            result={'id':1,'question':'Build a React calculator.','language':'react','status':'failed','code':'const value="\x00";',
                'output':'','error':'\x1b[31mGenuine Vite failure\x1b[0m\x00','screenshot_paths':[],
                'verification':{'required':True,'status':'failed','tests':[{'name':'Click Add','status':'failed','expected_text':['5'],'observed_text':'23\x00','error':'No expected result'}]}}
            workflow=SimpleNamespace(id=3,output_name='diagnostic',instructions='',results=[result])
            upload=SimpleNamespace(file_type='docx',file_path=str(source),original_filename='manual.docx')
            with patch('app.services.assignment_workflow.settings.REPORT_DIR',str(root/'reports')):
                path,_=build_word_report(upload,workflow)
            text='\n'.join(p.text for p in Document(path).paragraphs)
            self.assertIn('Genuine Vite failure',text)
            self.assertIn('Observed page text: 23',text)
            self.assertNotIn('\x1b',text)
            self.assertIn('\x1b',result['error'])

    def test_original_word_tables_headers_and_images_survive_composition(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "output.png"
            Image.new("RGB", (120, 60), "blue").save(image)
            source = Document()
            source.add_heading("Original laboratory manual", 0)
            source.sections[0].header.paragraphs[0].text = "Original institution header"
            table = source.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Question in a table"
            table.cell(0, 1).text = "Calculate factorial of 5."
            source.add_picture(str(image))
            source_path = root / "manual.docx"
            source.save(source_path)
            upload = SimpleNamespace(file_type="docx", file_path=str(source_path), original_filename="manual.docx")
            self.assertIn("Calculate factorial of 5.", document_text(upload))
            # Original-document preservation belongs to the explicit full-report mode.
            workflow = SimpleNamespace(id=1, mode="theory_and_code", output_name="completed_lab", instructions="Use functions.", results=[{
                "id": 1, "question": "Calculate factorial of 5.", "language": "python", "answer": "Multiply the integers from 1 through 5.",
                "code": "print(120)", "output": "120", "error": "", "screenshot_paths": [str(image)],
            }])
            with patch("app.services.assignment_workflow.settings.REPORT_DIR", str(root / "reports")):
                report_path, filename = build_word_report(upload, workflow)
            report = Document(report_path)
            self.assertEqual(filename, "completed_lab.docx")
            self.assertEqual(report.tables[0].cell(0, 0).text, "Question in a table")
            self.assertEqual(report.sections[0].header.paragraphs[0].text, "Original institution header")
            self.assertEqual(len(report.inline_shapes), 2)
            self.assertEqual(report.paragraphs[0].text, "Original laboratory manual")
            appended = "\n".join(paragraph.text for paragraph in report.paragraphs[1:])
            self.assertIn("Program 1: Calculate factorial of 5.", appended)
            self.assertNotIn("Multiply the integers", appended)
            self.assertNotIn("print(120)", appended)
            self.assertNotIn("Captured output", appended)
            self.assertNotIn("Additional instructions", appended)
            # The original manual needs one separator before the first program.
            page_breaks = report.element.xpath('.//w:br[@w:type="page"]')
            self.assertEqual(len(page_breaks), 1)
            with ZipFile(report_path) as package:
                self.assertTrue(any(name.startswith("word/media/") for name in package.namelist()))

    def test_original_pdf_pages_are_included_in_word(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / "manual.pdf"
            Image.new("RGB", (600, 800), "white").save(source, "PDF", append_images=[Image.new("RGB", (600, 800), "lightblue")], save_all=True)
            upload = SimpleNamespace(file_type="pdf", file_path=str(source), original_filename="manual.pdf")
            workflow = SimpleNamespace(id=2, output_name="pdf_report", instructions="", results=[])
            with patch("app.services.assignment_workflow.settings.REPORT_DIR", str(root / "reports")):
                report_path, _ = build_word_report(upload, workflow)
            self.assertEqual(len(Document(report_path).inline_shapes), 2)

    def test_generated_python_cannot_read_system_paths_or_import_process_tools(self):
        for code in ["import subprocess\nsubprocess.run(['whoami'])", "open('../secret', 'r')", "print(eval('1 + 1'))", "import os\nprint(os.environ)"]:
            with self.subTest(code=code), self.assertRaises(ValueError):
                validate_generated_code(code, "python")
        validate_generated_code("import math\nprint(math.factorial(5))", "python")


if __name__ == "__main__":
    unittest.main()
