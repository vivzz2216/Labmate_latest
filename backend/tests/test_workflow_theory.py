"""Theory mode retains prose while programming answers remain screenshots-only."""

import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from docx import Document
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import AssignmentWorkflow, Report, Upload, User, UserProfile
from app.services.assignment_workflow import process_assignment
from app.services.lab_question_filter import filter_lab_tasks, is_handwritten_question


class TheoryFilterTests(unittest.TestCase):
    def test_mode_keeps_theory_and_code_but_never_handwritten_or_headings(self):
        tasks = [
            {"text": "Explain the concept of inheritance in Java.", "language": "theory"},
            {"text": "Write a Python program to calculate factorial.", "language": "python"},
            {"text": "Sketch the flowchart by hand for a sorting algorithm.", "language": "theory"},
            {"text": "Theory Questions", "language": "theory"},
            {"text": "Explain inheritance and write a Java program to demonstrate it.", "language": "theory"},
        ]
        self.assertEqual([item["text"] for item in filter_lab_tasks(tasks, mode="code_only")], [tasks[1]["text"]])
        mixed = filter_lab_tasks(tasks, mode="theory_and_code")
        self.assertEqual([item["text"] for item in mixed], [tasks[0]["text"], tasks[1]["text"], tasks[4]["text"]])
        self.assertEqual([item["is_theory"] for item in mixed], [True, False, False])
        self.assertEqual(mixed[-1]["language"], "java")
        for phrase in ("handwritten", "hand-written", "by hand", "draw neatly", "sketch"):
            self.assertTrue(is_handwritten_question(f"Please {phrase} this answer"))


class TheoryProcessingTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)
        source = Document()
        source.add_paragraph("Original assignment remains in this document.")
        self.source_path = self.root / "manual.docx"
        source.save(self.source_path)
        self.image_path = self.root / "screen.png"
        Image.new("RGB", (200, 100), "white").save(self.image_path)
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.factory = sessionmaker(bind=self.engine)
        with self.factory() as db:
            user = User(email="actual@example.test", name="Asha Rao")
            db.add(user)
            db.flush()
            db.add(UserProfile(user_id=user.id, course="Computer Engineering", institution="Real University",
                               profile_metadata={"usn": "USN-9876"}))
            upload = Upload(user_id=user.id, filename="manual.docx", original_filename="manual.docx",
                            file_path=str(self.source_path), file_type="docx", file_size=self.source_path.stat().st_size)
            db.add(upload)
            db.flush()
            workflow = AssignmentWorkflow(upload_id=upload.id, status="processing", stage="generate", mode="theory_and_code",
                screenshot_style="style_2", questions=[
                    {"id": 1, "text": "Explain inheritance in Java.", "language": "theory", "is_theory": True},
                    {"id": 2, "text": "Write a Python program to print a greeting.", "language": "python"},
                ], results=[], language="auto", output_name="mixed_report")
            db.add(workflow)
            db.commit()
            self.workflow_id = workflow.id

    async def asyncTearDown(self):
        self.engine.dispose()
        self.folder.cleanup()

    async def test_theory_skips_runtime_and_code_is_not_written_to_word(self):
        generated = AsyncMock(side_effect=[
            {"answer": "Inheritance lets a subclass reuse behavior."},
            {"answer": "Greeting example", "code": "print('Asha Rao')", "language": "python", "stdin": ""},
        ])
        with patch("app.services.assignment_workflow.SessionLocal", self.factory), \
             patch("app.services.assignment_workflow.generation_service.generate_json", generated), \
             patch("app.services.assignment_workflow.execute_solution", AsyncMock(return_value=(True, "Asha Rao", "", []))) as execute, \
             patch("app.services.assignment_workflow.screenshot_service.generate_screenshot", AsyncMock(return_value=(True, str(self.image_path), 200, 100))) as capture, \
             patch("app.services.assignment_workflow.settings.REPORT_DIR", str(self.root / "reports")):
            await process_assignment(self.workflow_id)
        with self.factory() as db:
            workflow = db.get(AssignmentWorkflow, self.workflow_id)
            self.assertEqual(workflow.status, "completed", workflow.error)
            self.assertEqual(len(workflow.results), 2)
            self.assertEqual(workflow.results[0]["screenshot_paths"], [])
            self.assertEqual(execute.await_count, 1)
            self.assertEqual(capture.await_count, 2)
            prompt = generated.await_args_list[1].args[1]
            for value in ("Asha Rao", "USN-9876", "Computer Engineering", "Real University"):
                self.assertIn(value, prompt)
            report = db.get(Report, workflow.report_id)
            document = Document(report.file_path)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Original assignment remains", text)
            self.assertIn("Inheritance lets a subclass reuse behavior.", text)
            self.assertNotIn("print('Asha Rao')", text)
            self.assertNotIn("Greeting example", text)
            self.assertEqual(len(document.inline_shapes), 2)

    async def test_generated_profile_uses_requested_fallback_details(self):
        with self.factory() as db:
            profile = db.query(UserProfile).first()
            profile.course = "Computer Science and Engineering"
            profile.institution = "Government Engineering College"
            profile.profile_metadata = {"auto_generated": True}
            workflow = db.get(AssignmentWorkflow, self.workflow_id)
            workflow.questions = [workflow.questions[0]]
            db.commit()
        generated = AsyncMock(return_value={"answer": "Inheritance shares behavior."})
        with patch("app.services.assignment_workflow.SessionLocal", self.factory), \
             patch("app.services.assignment_workflow.generation_service.generate_json", generated), \
             patch("app.services.assignment_workflow.settings.REPORT_DIR", str(self.root / "reports")):
            await process_assignment(self.workflow_id)
        prompt = generated.await_args.args[1]
        for value in ("Asha Rao", "USN-2024-001", "Computer Science & Engineering", "Engineering College"):
            self.assertIn(value, prompt)
