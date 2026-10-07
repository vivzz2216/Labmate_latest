import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

import httpx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image

from app.services.docx_layout import embed_screenshot, save_document_atomic
from app.services.lab_question_filter import filter_lab_tasks, programming_lines
from app.services.parser_service import parser_service
from app.services.retry import RetryExhausted, RetryPolicy, retry_request
from app.services.runtime_engine import detect_cpp, java_stdin, runtime_engine
from app.services.screenshot_service import screenshot_service


class RetryTests(unittest.IsolatedAsyncioTestCase):
    async def test_503_backoff_recovers_without_dropping_request(self):
        operation = AsyncMock(side_effect=[httpx.Response(503), httpx.Response(503), httpx.Response(200)])
        sleep = AsyncMock()
        notify = AsyncMock()
        response = await retry_request(operation, RetryPolicy(4, 1, 10), notify, sleep, lambda low, high: high / 2)
        self.assertEqual(response.status_code, 200)
        self.assertEqual([call.args[0] for call in sleep.await_args_list], [.5, 1.0])
        self.assertEqual(operation.await_count, 3)
        self.assertEqual(notify.await_count, 2)

    async def test_retry_after_network_failure_and_permanent_auth_error(self):
        operation = AsyncMock(side_effect=[httpx.ConnectError("network"), httpx.Response(503, headers={"Retry-After": "7"}), httpx.Response(200)])
        sleep = AsyncMock()
        await retry_request(operation, RetryPolicy(3, 1, 4), sleep=sleep, uniform=lambda low, high: high)
        self.assertEqual([call.args[0] for call in sleep.await_args_list], [1, 7])
        denied = AsyncMock(return_value=httpx.Response(401))
        self.assertEqual((await retry_request(denied, sleep=sleep)).status_code, 401)
        self.assertEqual(denied.await_count, 1)

    async def test_exhaustion_and_cancellation_are_explicit(self):
        with self.assertRaises(RetryExhausted):
            await retry_request(AsyncMock(return_value=httpx.Response(503)), RetryPolicy(2, 0, 0), sleep=AsyncMock())
        with self.assertRaises(asyncio.CancelledError):
            await retry_request(AsyncMock(side_effect=asyncio.CancelledError()))


class FilteringTests(unittest.TestCase):
    def test_programming_only_and_theory_labels_are_rejected(self):
        tasks = [{"question_text": text} for text in ["Define inheritance.", "Explain the output of a Python program.", "Q1. Explain how to write a Python program.", "Objective: Write a Python program.", "Viva: Create a function.", "Concept questions", "Study concepts through Python code.", "Write a Python program to compute factorial.", "Implement inheritance using Java classes."]]
        kept = filter_lab_tasks(tasks)
        self.assertEqual([task["id"] for task in kept], [1, 2])
        self.assertTrue(all(task["question_text"].startswith(("Write", "Implement")) for task in kept))

    def test_sections_stay_excluded_until_next_lab_heading(self):
        source = ["Experiment 1", "Write a program to reverse a string.", "Viva Questions", "1. Write a program to explain this concept.", "Define a stack.", "Experiment 2", "Implement a stack using an array."]
        cleaned = programming_lines(source)
        self.assertNotIn(source[3], cleaned)
        self.assertIn(source[-1], cleaned)

    def test_ai_parser_cannot_bypass_the_filter(self):
        with patch.object(parser_service, "_extract_tasks_with_openai", return_value=[
            {"question_text": "Explain Java classes.", "code_snippet": "class A {}"},
            {"question_text": "Write a Java program to add two numbers.", "code_snippet": "class Main {}"},
        ]):
            result = parser_service._extract_tasks_from_lines(["Write a Java program to add two numbers."])
        self.assertEqual(len(result), 1)
        self.assertIn("add two numbers", result[0]["question_text"])


class ComposerLayoutTests(unittest.TestCase):
    def test_inline_images_fit_section_and_keep_original_styles(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = root / "wide.png"
            Image.new("RGB", (3000, 400), "blue").save(image, dpi=(96, 96))
            doc = Document()
            doc.styles["Normal"].font.name = "Arial"
            doc.styles["Normal"].font.size = Pt(13)
            doc.sections[0].left_margin = Inches(1.8)
            doc.sections[0].header.paragraphs[0].text = "Original header"
            doc.add_table(rows=1, cols=1).cell(0, 0).text = "Original table"
            shape = embed_screenshot(doc, image, "Actual captured output")
            section = doc.sections[0]
            self.assertLessEqual(shape.width, section.page_width - section.left_margin - section.right_margin)
            self.assertAlmostEqual(shape.width / shape.height, 7.5, places=3)
            image_paragraph = next(p for p in doc.paragraphs if p._p.xpath('.//w:drawing'))
            self.assertEqual(image_paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            path = root / "report.docx"
            save_document_atomic(doc, path)
            reopened = Document(path)
            self.assertEqual(reopened.styles["Normal"].font.name, "Arial")
            self.assertEqual(reopened.styles["Normal"].font.size, Pt(13))
            self.assertEqual(reopened.tables[0].cell(0, 0).text, "Original table")
            self.assertEqual(reopened.sections[0].header.paragraphs[0].text, "Original header")
            self.assertFalse(path.with_suffix('.partial.docx').exists())

    def test_tall_screenshot_is_sliced_without_losing_pixels(self):
        from io import BytesIO
        from PIL import ImageChops

        with tempfile.TemporaryDirectory() as folder:
            image_path = Path(folder) / "tall.png"
            original = Image.new("RGB", (1000, 3900), "white")
            for y in range(0, 3900, 75):
                for x in range(1000):
                    original.putpixel((x, y), (y % 255, x % 255, 60))
            original.save(image_path, dpi=(96, 96))
            doc = Document()
            embed_screenshot(doc, image_path)
            self.assertGreater(len(doc.inline_shapes), 1)
            pieces = [Image.open(BytesIO(doc.part.related_parts[
                shape._inline.graphic.graphicData.pic.blipFill.blip.embed].blob)).convert("RGB")
                for shape in doc.inline_shapes]
            self.assertTrue(all(piece.width == 1000 for piece in pieces))
            self.assertEqual(sum(piece.height for piece in pieces), original.height)
            restored = Image.new("RGB", original.size)
            top = 0
            for piece in pieces:
                restored.paste(piece, (0, top))
                top += piece.height
            self.assertIsNone(ImageChops.difference(original, restored).getbbox())


class RuntimeTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.root = Path(self.folder.name)
        self.patches = [patch("app.services.runtime_engine.settings.REACT_TEMP_DIR", str(self.root / "runtime")), patch("app.services.runtime_engine.settings.SCREENSHOT_DIR", str(self.root / "screenshots"))]
        for item in self.patches:
            item.start()

    async def asyncTearDown(self):
        for item in self.patches:
            item.stop()
        self.folder.cleanup()

    async def test_python_stdin_timeout_output_limit_and_secret_isolation(self):
        result = await runtime_engine.execute("n=int(input())\nprint(n*n)", "python", stdin="7\n")
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.output.splitlines(), ["7", "49"])
        with patch("app.services.runtime_engine.settings.EXECUTION_TIMEOUT", .25):
            result = await runtime_engine.execute("while True: pass", "python")
        self.assertEqual(result.exit_code, 124)
        with patch("app.services.runtime_engine.settings.EXECUTION_MAX_OUTPUT_BYTES", 512):
            result = await runtime_engine.execute("print('x'*2000)", "python")
        self.assertEqual(result.exit_code, 125)
        from app.services.runtime_engine import runtime_environment
        with patch.dict("os.environ", {"GEMINI_API_KEY": "sensitive-test", "GROQ_API_KEY": "groq-sensitive-test", "DATABASE_URL": "private-test"}):
            environment = runtime_environment(self.root)
        self.assertNotIn("GEMINI_API_KEY", environment)
        self.assertNotIn("GROQ_API_KEY", environment)
        self.assertNotIn("DATABASE_URL", environment)

    async def test_java_alias_scanner_and_concurrent_class_names(self):
        code = 'import java.util.Scanner; public class Main { public static void main(String[] args) { Scanner in=new Scanner(System.in); int n=in.nextInt(); String s=in.next(); System.out.println(n+":"+s); } }'
        first, second = await asyncio.gather(runtime_engine.execute(code, "java", stdin="7\none\n"), runtime_engine.execute(code, "java", stdin="9\ntwo\n"))
        self.assertTrue(first.success, first.error)
        self.assertTrue(second.success, second.error)
        self.assertEqual(first.output.strip(), "7:one")
        self.assertEqual(second.output.strip(), "9:two")
        self.assertEqual(java_stdin('in.nextInt(); in.nextLine(); in.nextLine();'), "5\nsample\n")

    async def test_c_and_cpp_auto_detection_and_compilation_errors(self):
        code_c = '#include <stdio.h>\nint main(){int n;scanf("%d",&n);printf("%d\\n",n*2);return 0;}'
        result = await runtime_engine.execute(code_c, "c", stdin="8\n")
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.output.strip(), "16")
        code_cpp = '#include <iostream>\nint main(){int n;std::cin>>n;std::cout<<n*n;}'
        self.assertTrue(detect_cpp(code_cpp))
        self.assertFalse(detect_cpp('#include <string.h>\nint main(){}'))
        result = await runtime_engine.execute(code_cpp, "c", stdin="6\n")
        self.assertTrue(result.success, result.error)
        self.assertEqual(result.output.strip(), "36")
        result = await runtime_engine.execute("int main(){ invalid syntax }", "c")
        self.assertFalse(result.success)
        self.assertIn("Compilation error", result.error)

    async def test_html_runs_real_dom_and_captures_image(self):
        result = await runtime_engine.execute('<!doctype html><h1>Actual browser output</h1><script>document.body.dataset.ready="yes"</script>', "html")
        self.assertTrue(result.success, result.error)
        self.assertIn("Actual browser output", result.output)
        self.assertTrue(Path(result.screenshots["/"]).is_file())
        with Image.open(result.screenshots["/"]) as image:
            self.assertGreater(image.width, 1000)

    async def test_protected_route_401_is_valid_captured_output(self):
        body = b"Authentication required"

        async def serve(reader, writer):
            await reader.readuntil(b"\r\n\r\n")
            writer.write(b"HTTP/1.1 401 Unauthorized\r\nContent-Type: text/plain\r\n"
                         + f"Content-Length: {len(body)}\r\nConnection: close\r\n\r\n".encode()
                         + body)
            await writer.drain()
            writer.close()

        server = await asyncio.start_server(serve, "127.0.0.1", 0)
        try:
            async with server:
                port = server.sockets[0].getsockname()[1]
                result = await runtime_engine._capture(url=f"http://127.0.0.1:{port}", routes=["/profile"])
            self.assertTrue(result.success, result.error)
            self.assertIn("Authentication required", result.output)
            self.assertTrue(Path(result.screenshots["/profile"]).is_file())
        finally:
            server.close()
            await server.wait_closed()

    async def test_theme_capture_contains_actual_code_without_dummy_output(self):
        for theme in ("idle", "notepad", "codeblocks"):
            rendered = await screenshot_service._render_template("print(7)", "", theme, filename="Actual.java")
            self.assertNotIn("Hello, World!", rendered)
            self.assertIn("print(7)", rendered)
            success, path, width, height = await screenshot_service.generate_screenshot('print("verified")', "verified", theme, username="Runtime test")
            self.assertTrue(success, path)
            self.assertGreater(width, 0)
            self.assertGreater(height, 0)

    async def test_node_ports_and_react_browser_capture(self):
        node_code = "const http=require('node:http');http.createServer((req,res)=>res.end('Actual Node output')).listen(3000);"
        first, second = await asyncio.gather(runtime_engine.execute(node_code, "node"), runtime_engine.execute(node_code, "node"))
        self.assertTrue(first.success, first.error)
        self.assertTrue(second.success, second.error)
        self.assertIn("Actual Node output", first.output)
        self.assertNotEqual(first.screenshots["/"], second.screenshots["/"])
        result = await runtime_engine.execute("import React from 'react';export default function App(){return <h1>Actual React output</h1>}", "react")
        self.assertTrue(result.success, result.error)
        self.assertIn("Actual React output", result.output)
        self.assertTrue(Path(result.screenshots["/"]).exists())

    async def test_compatibility_browser_capture_blocks_response_scripts(self):
        response = '<h1>Actual response</h1><script>document.body.innerHTML="";document.body.style.background="rgb(255,0,0)"</script>'
        success, path, width, height = await screenshot_service.generate_browser_screenshot(response, 7, "http://localhost:3000")
        self.assertTrue(success, path)
        self.assertEqual(width, 1366)
        self.assertGreaterEqual(height, 768)
        with Image.open(path) as image:
            self.assertNotEqual(image.convert("RGB").getpixel((100, 500)), (255, 0, 0))


if __name__ == "__main__":
    unittest.main()
