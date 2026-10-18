"""Real browser -> workflow result -> Word -> owned HTTP download.

Generation is deterministic here; provider availability is tested separately.
"""
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base, get_db
from app.middleware.auth import get_current_user
from app.models import AssignmentWorkflow, Report, Upload, User
from app.services.assignment_workflow import process_assignment


@pytest.mark.asyncio
@pytest.mark.parametrize('wrong',[False,True],ids=['verified-add','wrong-sum-diagnostic'])
async def test_actual_add_browser_word_preview_and_download(tmp_path,wrong):
    objective='Build an HTML calculator to add two numbers when Add is clicked.'
    source=Document();source.add_paragraph(objective);source_path=tmp_path/'manual.docx';source.save(source_path)
    engine=create_engine('sqlite:///'+str(tmp_path/'state.db'),connect_args={'check_same_thread':False})
    Base.metadata.create_all(engine);factory=sessionmaker(bind=engine)
    with factory() as db:
        user=User(email='integration@example.com',name='Integration Student');db.add(user);db.flush();user_id=user.id
        upload=Upload(user_id=user.id,filename='manual.docx',original_filename='manual.docx',file_path=str(source_path),file_type='docx',file_size=source_path.stat().st_size)
        db.add(upload);db.flush();upload_id=upload.id
        workflow=AssignmentWorkflow(upload_id=upload.id,status='processing',stage='generate',language='auto',questions=[{'id':7,'text':objective,'language':'html'}],results=[],output_name='add_test')
        db.add(workflow);db.commit();workflow_id=workflow.id
    operation='a.value+b.value' if wrong else 'Number(a.value)+Number(b.value)'
    html=f'''<main><label for="a">First</label><input id="a"><label for="b">Second</label><input id="b"><button id="add">Add</button><output id="result"></output></main><script>document.getElementById('add').onclick=()=>document.getElementById('result').textContent={operation}</script>'''
    response={'answer':'Add the entered numbers.','entrypoint':'index.html','files':[{'path':'index.html','content':html}],
        'browser_tests':[{'name':'Add 2 and 3','requirement_ids':['calculation'],'viewport':{'width':1366,'height':768},'steps':[
            {'action':'fill','locator_type':'label','locator':'First','role':'','value':'2'},
            {'action':'fill','locator_type':'label','locator':'Second','role':'','value':'3'},
            {'action':'click','locator_type':'role','locator':'Add','role':'button','value':''},
            {'action':'assert_visible','locator_type':'text','locator':'5','role':'','value':''}]}]}
    with patch('app.services.assignment_workflow.SessionLocal',factory),patch('app.services.assignment_workflow.generation_service.generate_json',AsyncMock(side_effect=lambda *a,**kw:deepcopy(response))),patch('app.services.assignment_workflow.settings.REPORT_DIR',str(tmp_path/'reports')),patch('app.services.assignment_workflow.settings.SCREENSHOT_DIR',str(tmp_path/'screenshots')),patch('app.services.assignment_workflow.settings.REACT_TEMP_DIR',str(tmp_path/'runtime')):
        await process_assignment(workflow_id)
    with factory() as db:
        workflow=db.get(AssignmentWorkflow,workflow_id)
        assert workflow.status==('failed' if wrong else 'completed'),workflow.error
        assert workflow.results[0]['id']==7
        assert workflow.report_id
        assert workflow.results[0]['verification']['status']==('failed' if wrong else 'passed')
        assert len(workflow.results[0]['screenshot_paths'])>=2
        report=db.get(Report,workflow.report_id)
        document=Document(report.file_path)
        assert document.paragraphs[0].text==objective
        if wrong:
            text='\n'.join(p.text for p in document.paragraphs)
            assert 'Expected visible text: 5' in text
            assert 'Observed page text:' in text and '23' in text
            assert 'Diagnostic report' in text
    from app.routers import workflows,download
    app=FastAPI();app.include_router(workflows.router,prefix='/api');app.include_router(download.router,prefix='/api')
    def database():
        with factory() as db: yield db
    app.dependency_overrides[get_db]=database
    app.dependency_overrides[get_current_user]=lambda:SimpleNamespace(id=user_id,email='integration@example.com',name='Integration Student')
    with TestClient(app) as client:
        result=client.get(f'/api/workflows/{upload_id}');assert result.status_code==200,result.text
        assert 'project_files' not in result.json()['results'][0]
        assert result.json()['results'][0]['project_file_names']==['index.html']
        downloaded=client.get(result.json()['report']['download_url']);assert downloaded.status_code==200
        with ZipFile(BytesIO(downloaded.content)) as package:
            images=[n for n in package.namelist() if n.startswith('word/media/')]
            assert len(images)>=2
            for name in images:
                with Image.open(BytesIO(package.read(name))) as image:image.verify()
        preview=client.get(f'/api/workflows/{upload_id}/preview');assert preview.status_code==200
        assert 'data:image/png;base64,' in preview.text
    engine.dispose()
