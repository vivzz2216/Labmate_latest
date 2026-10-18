"""Real runner and report gates with deterministic generation fixtures."""
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import AssignmentWorkflow, Report, Upload, User
from app.services.assignment_workflow import process_assignment
from app.services.evidence_planner import plan_experiment, select_evidence
from PIL import Image


@pytest.mark.asyncio
@pytest.mark.parametrize('case', ['wrong_output', 'missing_output', 'language_mismatch', 'gui_substitution'])
async def test_failed_requirement_produces_diagnostic_without_false_success(case, tmp_path):
    source = tmp_path/'manual.docx'
    doc = Document()
    doc.add_paragraph('UNRELATED EXPERIMENT MUST NOT ENTER GENERATION CONTEXT')
    doc.save(source)
    engine = create_engine('sqlite:///:memory:')
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    objective = 'Write a Python program to print the sum.\nExpected output: 5'
    requested = 'python'
    code, language = "print('23')", 'python'
    if case == 'missing_output': code = 'pass'
    if case == 'language_mismatch': language = 'node'
    if case == 'gui_substitution':
        requested = 'java'
        objective = 'Create a Java Swing calculator with a button'
    with factory() as db:
        user = User(email='audit@example.com', name='Audit Student')
        db.add(user); db.flush()
        upload = Upload(user_id=user.id, filename='manual.docx', original_filename='manual.docx', file_path=str(source), file_type='docx', file_size=source.stat().st_size)
        db.add(upload); db.flush()
        workflow = AssignmentWorkflow(upload_id=upload.id, status='processing', stage='generate',
            questions=[{'id': 1, 'text': objective, 'language': requested}], results=[], language='auto', output_name=case)
        db.add(workflow); db.commit(); workflow_id = workflow.id
    generate = AsyncMock(return_value={'answer': 'Exercise implementation.', 'code': code, 'language': language, 'stdin': '', 'browser_tests': []})
    with patch('app.services.assignment_workflow.SessionLocal', factory), \
         patch('app.services.assignment_workflow.generation_service.generate_json', generate), \
         patch('app.services.assignment_workflow.settings.REPORT_DIR', str(tmp_path/'reports')), \
         patch('app.services.assignment_workflow.settings.SCREENSHOT_DIR', str(tmp_path/'screenshots')), \
         patch('app.services.assignment_workflow.settings.REACT_TEMP_DIR', str(tmp_path/'runtime')):
        await process_assignment(workflow_id)
    with factory() as db:
        workflow = db.get(AssignmentWorkflow, workflow_id)
        assert workflow.status == 'failed'
        assert workflow.results[0]['status'] == 'failed'
        assert workflow.report_id
        report = db.get(Report, workflow.report_id)
        assert Path(report.file_path).is_file()
        assert 'has not passed' in '\n'.join(p.text for p in Document(report.file_path).paragraphs)
    assert 'UNRELATED EXPERIMENT' not in generate.call_args.args[1]
    engine.dispose()


def test_intermediate_requirement_cannot_be_satisfied_by_one_duplicate_image(tmp_path):
    path = tmp_path/'state.png'
    Image.new('RGB', (320,200), 'white').save(path)
    plan = plan_experiment('Create an HTML form. Show invalid input and successful submission.', 'html')
    candidate = {'path': str(path), 'kind': 'browser'}
    selection = select_evidence(plan, [candidate, candidate])
    assert 'intermediate' in selection['missing_requirements']
