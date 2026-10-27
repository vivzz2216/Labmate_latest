"""Actual native execution and Word evidence through the assignment workflow."""
from unittest.mock import AsyncMock
from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app.database import Base
from app.models import User,Upload,AssignmentWorkflow,Report
from app.config import settings
from app.services import assignment_workflow as flow
from app.services.manual_sections import numbered_experiments
from app.routers.workflows import QuestionInput


def test_new_lab_extraction_and_contracts():
    cases=numbered_experiments('Experiment 1: SQL joins\nAim: Write SQL queries joining tables.\nExperiment 2: Colab notebook\nAim: Print a Python calculation.')
    assert [q['language'] for q in cases]==['sql','colab']
    for case in cases:assert QuestionInput(**{k:case[k] for k in ('id','text','language')}).language==case['language']


def test_dataset_context_supplies_names_and_headers_without_all_training_rows():
    brief=flow.dataset_brief({'teacher.csv':'x,target\n1,private-training-row\n'})
    assert 'teacher.csv' in brief and 'x,target' in brief
    assert 'private-training-row' not in brief


@pytest.mark.asyncio
@pytest.mark.parametrize('language,code,expected',[
    ('sql','CREATE TABLE t(v INT);INSERT INTO t VALUES(NULL),(10),(20);SELECT AVG(v) AS average FROM t;',True),
    ('colab','print(7//2)',True),
    ('python','print(7//2)',True),
    ('colab','import pandas as pd\nprint(pd.read_csv("teacher.csv")["value"].mean())',True),
    ('sql','SELECT missing FROM missing_table;',False),
    ('colab','raise ValueError("deliberate notebook failure")',False),
])
async def test_lab_execution_status_and_word_evidence(tmp_path,monkeypatch,language,code,expected):
    monkeypatch.setattr(settings,'SCREENSHOT_DIR',str(tmp_path/'screenshots'))
    monkeypatch.setattr(settings,'REPORT_DIR',str(tmp_path/'reports'))
    monkeypatch.setattr(settings,'BILLING_ENABLED',False)
    monkeypatch.setattr(settings,'EXECUTION_ALLOWED_EMAILS','audit@example.com')
    monkeypatch.setattr(settings,'LLM_PROVIDER','glm')
    source=tmp_path/'manual.docx';doc=Document();doc.add_paragraph('Pipeline regression manual');doc.save(source)
    engine=create_engine('sqlite:///'+(tmp_path/'state.db').as_posix(),connect_args={'check_same_thread':False})
    Base.metadata.create_all(engine);factory=sessionmaker(bind=engine)
    with factory() as db:
        user=User(email='audit@example.com',name='Audit');db.add(user);db.flush()
        upload=Upload(user_id=user.id,filename=source.name,original_filename=source.name,file_path=str(source),file_type='docx',file_size=source.stat().st_size)
        db.add(upload);db.flush()
        question={'id':1,'language':language,'text':'Write SQL queries using SQLite and show the result.' if language=='sql' else 'Use a Colab notebook to print 7 divided by 2 using floor division.'}
        if 'read_csv' in code:
            question['text']='Use a Colab notebook to compute the mean of value in the supplied teacher.csv.'
            question['datasets']={'teacher.csv':'value\n10\n20\n'}
            question['colab_theme']='light'
        wf=AssignmentWorkflow(upload_id=upload.id,status='processing',stage='generate',mode='code_only',language='auto',questions=[question],results=[],output_name='lab_regression')
        db.add(wf);db.commit();wid=wf.id
    monkeypatch.setattr(flow,'SessionLocal',factory)
    if language in {'colab','python'}:
        monkeypatch.setattr(flow,'capture_source_evidence',AsyncMock(side_effect=AssertionError('Colab must not use a generic source renderer')))
    monkeypatch.setattr(flow.generation_service,'generate_json',AsyncMock(return_value={'language':language,'answer':'Complete source and execution evidence.','code':code,'stdin':'','browser_tests':[]}))
    await flow.process_assignment(wid)
    with factory() as db:
        wf=db.get(AssignmentWorkflow,wid)
        assert wf.status==('completed' if expected else 'failed'),(wf.status,wf.error,wf.results)
        assert wf.report_id, (wf.status,wf.error)
        result=wf.results[0]
        if 'read_csv' in code:assert result['output']=='15.0',result
        if 'read_csv' in code:assert result['verification']['theme']=='light'
        assert result['verification']['status']==('passed' if expected else 'failed')
        assert len(result['screenshot_paths'])>=2,result.get('error')
        report=db.get(Report,wf.report_id)
        path=Path(report.file_path)
        word=Document(path)
        assert len(word.inline_shapes)>=2
        if not expected:
            assert 'diagnostic report' in '\n'.join(p.text for p in word.paragraphs).lower()
    engine.dispose()
