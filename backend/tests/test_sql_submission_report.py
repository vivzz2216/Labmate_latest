from types import SimpleNamespace
from pathlib import Path
from PIL import Image
from docx import Document
from app.config import settings
from app.services.sql_report import clean_report, post_questions, report_number, answer_question
from app.services.sql_result_checks import independent_checks


def test_report_labels_and_post_questions():
    assert 'roles' in answer_question('What is role based access control?')
    assert report_number('SQL_Exp_8.docx', 'Experiment 7: Triggers')==8
    assert post_questions('Post Experimental Exercise\nA. Questions:\n1. What is DDL?\n2. Why no rollback when we use DDL? 3. What is the execution pattern of an SQL query?\nB. Conclusion:\nWrite the conclusion.') == ['What is DDL?','Why no rollback when we use DDL?','What is the execution pattern of an SQL query?']


def test_report_uses_receipt_counts_not_stale_model_description(tmp_path, monkeypatch):
    monkeypatch.setattr(settings,'REPORT_DIR',str(tmp_path))
    image=tmp_path/'capture.png';Image.new('RGB',(1440,900),'white').save(image)
    source='Experiment 7: Triggers\nInvalid manual date 2020-10-012\n**Key points**\nPost Experimental Exercise\nA. Questions:\nWhat are the conditions under which we should avoid using Triggers\nB. Conclusion:\nWrite what was performed.'
    step={'success':True,'index':1,'statement':'SELECT * FROM WorkCenterStats','columns':['totalCapacity'],'rows':[{'totalCapacity':300}]}
    result={'id':7,'language':'sql','status':'completed','question':'Experiment 7: Triggers\nAim: Perform triggers.',
            'answer':'Three tables: a,b,c,d,e. The total is 200.', 'verification':{'verified':True,'engine':'MySQL 8','statements':[step],
            'evidence_frames':[{'path':str(image),'kind':'database','caption':'Query: WorkCenterStats, 1 row'}]}}
    path,name=clean_report(SimpleNamespace(original_filename='SQL_Exp_8.docx'),SimpleNamespace(results=[result],output_name='triggers',id=1),source)
    doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
    assert text.startswith('Experiment 8 Triggers')
    assert 'Conclusion' in text and 'avoid triggers' in text.lower()
    for bad in ('**','not_specified','[PASS]','SQL Server examples','2020-10-012','total is 200','Write what was performed'):
        assert bad not in text
    assert doc.tables[0].rows[1].cells[2].text=='1'
    assert len(doc.inline_shapes)==1
    assert not doc.styles['Title']._element.xpath('.//w:pBdr')
    assert not doc.element.xpath('.//w:br[@w:type="page"]')


def test_join_verification_compares_values_and_duplicates():
    tables={'Orders':{'rows':[{'id':1,'emp':1},{'id':2,'emp':1}]},'Employees':{'rows':[{'id':1,'name':'A'},{'id':2,'name':'B'}]}}
    step={'index':1,'success':True,'statement':'SELECT o.id, e.name FROM Orders o RIGHT JOIN Employees e ON o.emp=e.id',
          'columns':['id','name'],'rows':[{'id':1,'name':'A'},{'id':2,'name':'A'},{'id':None,'name':'B'}],
          'schemas_snapshot':{'lab':{'tables':tables}},'active_schema':'lab'}
    assert independent_checks([step])[0]['passed']
    step['rows'][1]={'id':1,'name':'A'}
    assert not independent_checks([step])[0]['passed']


def test_procedure_captures_all_result_sets_and_definition():
    from app.services.sql_runtime import SQLSession
    session=SQLSession()
    try:
        r=session.execute_script('''DELIMITER $$
CREATE PROCEDURE example() BEGIN SELECT 'Mumbai' AS location; SELECT 'Delhi' AS location; END$$
DELIMITER ;
CALL example(); SHOW CREATE PROCEDURE example;''')
        assert r['success'], r['error']
        call=r['select_results'][0]
        assert call['rows']==[{'location':'Mumbai'},{'location':'Delhi'}]
        assert call['result_set_count']==2
        assert 'SELECT' in r['select_results'][1]['rows'][0]['Create Procedure']
    finally:session.close()
