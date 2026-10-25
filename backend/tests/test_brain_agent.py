import asyncio
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from docx import Document
from fastapi import HTTPException
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import User, Upload, AssignmentWorkflow
from app.services.brain_agent import make_plan, fingerprint, TOOLS, topic_choices
from app.services.generation_service import GenerationError
from app.services.source_evidence import capture_source_evidence, source_capture_plan
from app.services.evidence_planner import select_evidence
from app.services.assignment_workflow import build_word_report, process_assignment
from app.routers.workflows import clarify, ClarificationInput

QUESTION={'id':9,'language':'react','text':'Experiment - 9: Mini project. Build a webapp using HTML CSS JavaScript and React. Choose one topic: Restaurant; Music Store; Personal webpage.'}
DECISION={'decisions':[{'id':9,'topic':'Restaurant','reason':'A local menu and cart demonstrate React state without external services.','features':['Menu with categories','Filter by category','Add and remove cart items','Display the total']} ]}


class BrainTests(unittest.IsolatedAsyncioTestCase):
    async def test_topic_alternatives_are_read_from_unseen_manuals(self):
        self.assertEqual(topic_choices('Choose one topic from the following list:\n- Fitness Tracker\n- Quiz app\n8. Results:'),['Fitness Tracker','Quiz app'])
        self.assertEqual(topic_choices('Choose one topic: Restaurant; Music Store; Personal webpage.'),['Restaurant','Music Store','Personal webpage'])
        self.assertEqual(topic_choices('Suggested topics:\n1. Fitness Tracker\n2. Quiz app\n8. Results:'),['Fitness Tracker','Quiz app'])
    async def test_mini_project_chooses_one_source_topic_and_complete_image_policy(self):
        with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock(return_value=DECISION)):
            plan=await make_plan([QUESTION])
        self.assertEqual(plan['status'],'ready')
        task=plan['tasks'][0]
        self.assertEqual(task['topic'],'Restaurant')
        self.assertEqual(task['source_policy'],'complete')
        self.assertEqual(task['report_policy'],'images_only')
        self.assertFalse(task['include_description'])
        self.assertTrue(set(task['tools']).issubset(TOOLS))
        self.assertEqual(task['language'],'react')

    async def test_simple_task_requires_no_extra_provider_call(self):
        with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock()) as model:
            plan=await make_plan([{'id':1,'language':'python','text':'Print 6.'}])
        model.assert_not_awaited()
        self.assertEqual(plan['tasks'][0]['category'],'console')

    async def test_theory_keeps_its_existing_non_execution_path(self):
        with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock()) as model:
            plan=await make_plan([{'id':1,'language':'theory','is_theory':True,'text':'Explain mini-project planning.'}])
        model.assert_not_awaited()
        self.assertEqual(plan['tasks'][0]['tools'],['compose_report'])
        self.assertEqual(plan['tasks'][0]['report_policy'],'standard')

    async def test_framework_answer_applies_only_to_its_own_task(self):
        questions=[{'id':1,'language':'html','text':'Use Vue only.'},{'id':2,'language':'java','text':'Print 2.'}]
        plan=await make_plan(questions)
        plan['answers']={'framework-1':'Use React'}
        resumed=await make_plan(questions,previous=plan)
        self.assertEqual([t['language'] for t in resumed['tasks']],['react','java'])

    async def test_external_service_is_not_silently_replaced_with_fake_evidence(self):
        q={'id':1,'language':'react','text':'This app requires MySQL and real payments.'}
        plan=await make_plan([q])
        self.assertEqual(plan['status'],'awaiting_clarification')
        self.assertNotIn('Use a local demonstration',plan['questions'][0]['options'])
        plan['answers']={'service-1':'Revise the assignment'}
        with self.assertRaisesRegex(GenerationError,'not executed or verified'):
            await make_plan([q],previous=plan)

    async def test_reuses_saved_plan_but_changed_instructions_invalidate_it(self):
        with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock(return_value=DECISION)) as model:
            plan=await make_plan([QUESTION])
            self.assertEqual(await make_plan([QUESTION],previous=plan),plan)
            modified=await make_plan([QUESTION],'Include a description.',previous=plan)
        self.assertEqual(model.await_count,2)
        self.assertTrue(modified['tasks'][0]['include_description'])

    async def test_user_topic_is_preserved_and_unsupported_topic_is_rejected(self):
        with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock(return_value=DECISION)):
            plan=await make_plan([QUESTION],'Topic: Music Store')
        self.assertEqual(plan['tasks'][0]['topic'],'Music Store')
        bad={'decisions':[{**DECISION['decisions'][0],'topic':'Banking system'}]}
        with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock(return_value=bad)):
            with self.assertRaises(GenerationError):await make_plan([QUESTION])

    async def test_essential_question_pauses_before_generation_and_answer_survives(self):
        q={**QUESTION,'text':QUESTION['text']+' Ask me to select the topic.'}
        with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock(return_value=DECISION)) as model:
            plan=await make_plan([q])
            model.assert_not_awaited()
            self.assertEqual(plan['status'],'awaiting_clarification')
            self.assertLessEqual(len(plan['questions']),3)
            plan['answers']={'topic-9':'Personal webpage'}
            answered=await make_plan([q],previous=plan)
        self.assertEqual(answered['tasks'][0]['topic'],'Personal webpage')

    async def test_unknown_adapter_and_excessive_questions_never_execute(self):
        with self.assertRaises(GenerationError):
            await make_plan([{'id':1,'language':'python','text':'Build a tkinter GUI.'}])
        questions=[{'id':i,'language':'react','text':'Use Vue only.'} for i in range(1,5)]
        with self.assertRaises(GenerationError):await make_plan(questions)

    async def test_clarification_workflow_does_not_reserve_or_execute(self):
        with tempfile.TemporaryDirectory() as folder:
            engine=create_engine('sqlite:///'+folder.replace('\\','/')+'/test.db')
            Base.metadata.create_all(engine);factory=sessionmaker(bind=engine)
            with factory() as db:
                user=User(email='brain@example.com',name='Student');db.add(user);db.flush()
                upload=Upload(user_id=user.id,filename='x.docx',original_filename='x.docx',file_path='unused',file_type='docx',file_size=1);db.add(upload);db.flush()
                wf=AssignmentWorkflow(upload_id=upload.id,questions=[{**QUESTION,'text':QUESTION['text']+' Ask me which topic.'}],status='processing');db.add(wf);db.commit();wid=wf.id
            with patch('app.services.assignment_workflow.SessionLocal',factory),patch('app.services.assignment_workflow.settings.EXECUTION_ALLOWED_EMAILS','brain@example.com'),\
                 patch('app.services.billing_service.reserve_workflow_credits') as reserve,patch('app.services.assignment_workflow.execute_web_assignment',AsyncMock()) as execute:
                await process_assignment(wid)
            with factory() as db:self.assertEqual(db.get(AssignmentWorkflow,wid).status,'awaiting_clarification')
            reserve.assert_not_called();execute.assert_not_awaited();engine.dispose()

    async def test_every_source_line_and_css_is_captured_without_excerpts(self):
        files={'index.html':'<div id="root"></div>','src/main.jsx':'mount();','src/App.jsx':'\n'.join(f'const line{i}={i};' for i in range(100)),'style.css':'body{color:red;}'}
        async def capture(*args,**kwargs):return True,'image-'+str(capture.call_count),1128,800
        mocked=AsyncMock(return_value=(True,'unused.png',1128,800))
        with patch('app.services.source_evidence.screenshot_service.generate_screenshot',mocked):
            images,errors=await capture_source_evidence(files,'','index.html',theme='vscode',group='test',username='Student',language='react',screenshot_style='style_1',policy='complete')
        self.assertFalse(errors)
        self.assertEqual(set(x['source_file'] for x in images),set(files))
        lines={i for x in images if x['source_file']=='src/App.jsx' for a,b in x['source_ranges'] for i in range(a,b+1)}
        self.assertEqual(lines,set(range(1,101)))
        self.assertEqual(len([x for x in images if x['source_file']=='src/App.jsx']),4)

    async def test_missing_source_page_fails_complete_evidence(self):
        with tempfile.TemporaryDirectory() as folder:
            items=[]
            for i,(a,b) in enumerate([(1,32),(65,80)]):
                path=Path(folder)/f'{i}.png';Image.new('RGB',(200,100),(i*50,20,20)).save(path)
                items.append({'path':str(path),'kind':'source_preview','source_file':'App.jsx','source_ranges':[[a,b]]})
            result=select_evidence({'requirements':[],'complete_source_lines':{'App.jsx':80}},items)
            self.assertIn('source_lines:App.jsx',result['missing_requirements'])

    async def test_model_clarifications_must_be_grounded_and_never_request_secrets(self):
        missing='TBD: provide the dataset size.'
        q={'id':1,'language':'python','text':'Process a file. '+missing}
        request={'decisions':[],'clarifications':[{'id':'dataset-size','prompt':'How many rows should be processed?',
            'source_excerpt':missing,'options':[]}]}
        with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock(return_value=request)):
            plan=await make_plan([q])
        self.assertEqual(plan['status'],'awaiting_clarification')
        for prompt in ['What is your API key?','What is your password?']:
            bad={'decisions':[],'clarifications':[{**request['clarifications'][0],'prompt':prompt}]}
            with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock(return_value=bad)):
                with self.assertRaises(GenerationError):await make_plan([q])

    async def test_questions_stop_after_two_rounds_including_saved_answers(self):
        q={'id':1,'language':'python','text':'TBD: dataset and output format.'}
        previous=None
        for index in range(1,4):
            response={'decisions':[],'clarifications':[{'id':f'fact-{index}','prompt':'Which dataset size?',
                'source_excerpt':q['text'],'options':[]}]}
            with patch('app.services.brain_agent.generation_service.generate_json',AsyncMock(return_value=response)):
                if index==3:
                    with self.assertRaisesRegex(GenerationError,'two clarification rounds'):await make_plan([q],previous=previous)
                else:
                    previous=await make_plan([q],previous=previous)
                    self.assertEqual(previous['clarification_rounds'],index)
                    previous['answers'][f'fact-{index}']='100'

    async def test_images_only_report_contains_no_generated_source_or_test_transcript(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);original=Document();original.add_paragraph('Original mini-project manual');original.save(root/'manual.docx')
            Image.new('RGB',(500,300),'navy').save(root/'source.png')
            result={'id':9,'question':'Mini project','status':'completed','language':'react','answer':'Optional long explanation',
                    'code':'SECRET_SOURCE_MARKER','screenshot_paths':[str(root/'source.png')],'screenshot_captions':['Source App.jsx lines 1-32'],
                    'agent_decision':{'report_policy':'images_only','topic':'Restaurant','include_description':False}}
            with patch('app.services.assignment_workflow.settings.REPORT_DIR',str(root/'reports')):
                path,_=build_word_report(SimpleNamespace(file_path=str(root/'manual.docx'),file_type='docx'),SimpleNamespace(id=1,output_name='result',results=[result]))
            doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
            self.assertNotIn('SECRET_SOURCE_MARKER',text);self.assertNotIn('Optional long explanation',text)
            self.assertEqual(len(doc.inline_shapes),1)
