import unittest
from docx import Document
from app.services.lab_question_filter import (
    is_handwritten_question,
    is_theory_question,
    is_actionable_lab_question,
    filter_lab_tasks,
)
from app.services.docx_layout import embed_theory_block


class TheoryAndHandwrittenFilterTests(unittest.TestCase):
    def test_handwritten_detection(self):
        # Should detect explicit handwritten instructions
        self.assertTrue(is_handwritten_question("Write by hand the binary search tree traversal."))
        self.assertTrue(is_handwritten_question("This experiment must be handwritten on A4 sheets."))
        self.assertTrue(is_handwritten_question("Hand-written viva questions with diagrams."))
        self.assertTrue(is_handwritten_question("Draw by hand the state machine diagram."))
        self.assertTrue(is_handwritten_question("Draw neatly the architecture of 8086."))
        self.assertTrue(is_handwritten_question("Sketch the layout of the website."))

        # Normal questions should NOT be marked handwritten
        self.assertFalse(is_handwritten_question("Write a C program to implement quicksort."))
        self.assertFalse(is_handwritten_question("Explain the working of Dijkstra's algorithm."))
        self.assertFalse(is_handwritten_question("What is polymorphism in Object Oriented Programming?"))

    def test_theory_vs_code_filtering(self):
        tasks = [
            {"id": 1, "text": "Write a Python program to calculate factorial using recursion.", "language": "python"},
            {"id": 2, "text": "Explain the difference between shallow copy and deep copy in Python.", "language": "theory"},
            {"id": 3, "text": "Draw by hand the memory diagram of list slicing.", "language": "theory"},
            {"id": 4, "text": "Write a C++ program to implement a Binary Search Tree with insert and delete operations.", "language": "cpp"},
            {"id": 5, "text": "Handwritten notes: explain what a virtual destructor does in C++.", "language": "theory"},
        ]

        # Mode 1: code_only -> only programming exercises, zero theory
        code_only_results = filter_lab_tasks(tasks, mode="code_only")
        self.assertEqual(len(code_only_results), 2)
        self.assertEqual(code_only_results[0]["language"], "python")
        self.assertEqual(code_only_results[1]["language"], "cpp")
        self.assertFalse(code_only_results[0]["is_theory"])
        self.assertFalse(code_only_results[1]["is_theory"])

        # Mode 2: theory_and_code -> keeps programming exercises AND valid theory, skips all handwritten
        theory_and_code_results = filter_lab_tasks(tasks, mode="theory_and_code")
        self.assertEqual(len(theory_and_code_results), 3)
        self.assertEqual(theory_and_code_results[0]["text"], "Write a Python program to calculate factorial using recursion.")
        self.assertFalse(theory_and_code_results[0]["is_theory"])
        
        self.assertEqual(theory_and_code_results[1]["text"], "Explain the difference between shallow copy and deep copy in Python.")
        self.assertTrue(theory_and_code_results[1]["is_theory"])
        self.assertEqual(theory_and_code_results[1]["language"], "theory")

        self.assertEqual(theory_and_code_results[2]["text"], "Write a C++ program to implement a Binary Search Tree with insert and delete operations.")
        self.assertFalse(theory_and_code_results[2]["is_theory"])

    def test_embed_theory_block_layout(self):
        doc = Document()
        embed_theory_block(
            doc,
            question_text="What is the internal working of HashMap in Java?",
            answer_text="HashMap uses an array of Node buckets.\nWhen a key-value pair is inserted, hash(key) calculates the bucket index.\nIn Java 8+, collisions exceeding 8 nodes transform into balanced Red-Black Trees.",
            question_num=1,
        )
        
        # Verify document contains paragraphs and zero tables
        self.assertEqual(len(doc.tables), 0)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertTrue(any("Question 1:" in p for p in paragraphs))
        self.assertTrue(any("What is the internal working of HashMap" in p for p in paragraphs))
        self.assertTrue(any("HashMap uses an array of Node buckets" in p for p in paragraphs))


if __name__ == "__main__":
    unittest.main()
