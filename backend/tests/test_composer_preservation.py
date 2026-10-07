import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import Base
from app.models import Job, Screenshot, Upload
from app.services.composer_service import composer_service


class ComposerPreservationTests(unittest.IsolatedAsyncioTestCase):
    async def test_compose_preserves_source_and_respects_explicit_order(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = Document()
            source.styles["Normal"].font.name = "Arial"
            source.styles["Normal"].font.size = Pt(13)
            source.sections[0].left_margin = Inches(1.4)
            source.sections[0].header.paragraphs[0].text = "Original header"
            source.add_paragraph("Original assignment")
            source.add_table(rows=1, cols=1).cell(0, 0).text = "Original table"
            original = root / "manual.docx"
            source.save(original)
            original_bytes = original.read_bytes()
            image = root / "screenshot.png"
            Image.new("RGB", (1800, 900), "white").save(image)
            engine = create_engine("sqlite:///:memory:")
            Base.metadata.create_all(engine)
            with Session(engine) as db:
                upload = Upload(filename="manual.docx", original_filename="manual.docx", file_path=str(original), file_type="docx", file_size=1)
                db.add(upload); db.flush()
                jobs = [Job(upload_id=upload.id, task_id=i, question_text=f"Question {i}", code_snippet="print(1)", status="completed", output_text="1") for i in (1, 2)]
                db.add_all(jobs); db.flush()
                db.add_all([Screenshot(job_id=job.id, file_path=str(image), width=1800, height=900, file_size=image.stat().st_size) for job in jobs])
                db.commit()
                with patch("app.services.composer_service.settings.REPORT_DIR", str(root / "reports")):
                    result = await composer_service.compose_report(upload.id, [jobs[1].id, jobs[0].id], db)
            final = Document(result["report_path"])
            self.assertEqual(original.read_bytes(), original_bytes)
            self.assertEqual(final.styles["Normal"].font.name, "Arial")
            self.assertEqual(final.styles["Normal"].font.size, Pt(13))
            self.assertEqual(final.sections[0].left_margin, Inches(1.4))
            self.assertEqual(final.sections[0].header.paragraphs[0].text, "Original header")
            self.assertEqual(final.tables[0].cell(0, 0).text, "Original table")
            text = "\n".join(paragraph.text for paragraph in final.paragraphs)
            self.assertLess(text.index("Question 2"), text.index("Question 1"))
            self.assertNotIn("Programs and captured", text)
            self.assertNotIn("Description:", text)
            self.assertNotIn("Python code", text)
            self.assertEqual(len(final.inline_shapes), 2)
            engine.dispose()
