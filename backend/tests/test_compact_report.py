from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

from docx import Document
from PIL import Image
import pytest

from app.config import settings
from app.services import assignment_workflow as flow


def fixture_report(tmp_path,monkeypatch):
    monkeypatch.setattr(settings,'REPORT_DIR',str(tmp_path/'reports'))
    source=tmp_path/'manual.docx'
    doc=Document();doc.add_paragraph('Original manual cover and long instructions');doc.save(source)
    image=tmp_path/'evidence.png';Image.new('RGB',(600,300),'navy').save(image)
    upload=SimpleNamespace(file_path=str(source),file_type='docx',original_filename=source.name)
    # Report formatting must not send an AI request or reread the manual.
    monkeypatch.setattr(flow,'document_text',Mock(side_effect=AssertionError('Do not copy the original manual')))
    return source,image,upload


def test_code_reports_have_only_question_images_and_short_captions(tmp_path,monkeypatch):
    source,image,upload=fixture_report(tmp_path,monkeypatch)
    before=source.read_bytes()
    results=[]
    for number,language in enumerate(('python','java','html','colab','sql','cpp','c','react','node'),1):
        results.append(dict(id=number,language=language,status='completed',
            question=f'Experiment {number} {language} exercise\nLong detailed requirements',
            answer='A long explanation that must not appear.',code='RAW_SOURCE_MUST_NOT_APPEAR',
            project_files={'README.md':'Setup instructions must not appear.'},
            screenshot_paths=[str(image)],screenshot_captions=['Captured source and output.']))
    path,_=flow.build_word_report(upload,SimpleNamespace(id=1,mode='code_only',output_name='compact',results=results))
    doc=Document(path)
    assert len(doc.inline_shapes)==len(results)
    assert not doc.tables
    texts=[p.text for p in doc.paragraphs if p.text]
    assert texts==[text for r in results for text in (r['question'].splitlines()[0],'Captured source and output.')]
    assert source.read_bytes()==before
    for caption in [p for p in doc.paragraphs if p.text=='Captured source and output.']:
        assert caption._p.getprevious().xpath('.//w:drawing'),'Caption must follow its image'


def test_sql_default_uses_the_same_compact_structure(tmp_path,monkeypatch):
    _,image,upload=fixture_report(tmp_path,monkeypatch)
    result=dict(id=1,language='sql',status='completed',question='Experiment 1 Joins',
        answer='Do not add a conclusion or query summary table.',screenshot_paths=[str(image)],screenshot_captions=['MySQL join results.'])
    path,_=flow.build_word_report(upload,SimpleNamespace(id=1,output_name='sql',results=[result]))
    doc=Document(path)
    assert [p.text for p in doc.paragraphs if p.text]==['Experiment 1 Joins','MySQL join results.']
    assert not doc.tables
    assert doc.sections[0].page_width>doc.sections[0].page_height


def test_failed_report_keeps_a_short_diagnostic(tmp_path,monkeypatch):
    _,_,upload=fixture_report(tmp_path,monkeypatch)
    result=dict(id=1,language='java',status='failed',question='Compile a Java class',error='Compiler rejected the source.',screenshot_paths=[])
    path,_=flow.build_word_report(upload,SimpleNamespace(id=1,output_name='failed',results=[result]))
    texts=[p.text for p in Document(path).paragraphs if p.text]
    assert len(texts)==2
    assert 'has not passed verification' in texts[1] and 'Compiler rejected' in texts[1]


@pytest.mark.parametrize('missing',['absent','empty'])
def test_success_cannot_hide_missing_evidence(tmp_path,monkeypatch,missing):
    _,_,upload=fixture_report(tmp_path,monkeypatch)
    result=dict(id=1,language='python',status='completed',question='Print a number',
        screenshot_paths=[str(tmp_path/'missing.png')] if missing=='absent' else [])
    with pytest.raises(ValueError):
        flow.build_word_report(upload,SimpleNamespace(id=1,output_name='missing',results=[result]))
