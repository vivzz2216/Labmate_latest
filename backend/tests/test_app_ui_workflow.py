from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock,patch
from zipfile import ZipFile
import pytest
from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from app.database import Base,get_db
from app.middleware.auth import get_current_user
from app.models import AssignmentWorkflow,Report,Upload,User
from app.services.assignment_workflow import process_assignment
from fixtures.app_ui_cases import FLASK_FILES,FLASK_TESTS,SWING_SOURCE,GUI_TESTS


@pytest.mark.parametrize('language,wrong',[('python',False),('java',False),('java',True)],ids=['flask','swing','swing-diagnostic'])
@pytest.mark.asyncio
async def test_real_app_ui_workflow_word_preview_and_owned_download(tmp_path,language,wrong,monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, 'BILLING_ENABLED', True)
    objective=('Build a Flask calculator to add two numbers with server-side validation for invalid input.' if language=='python'
        else 'Build a Java Swing calculator to add two numbers, with validation for invalid input.')
    source=Document();source.add_paragraph(objective);manual=tmp_path/'manual.docx';source.save(manual)
    engine=create_engine('sqlite:///'+str(tmp_path/'state.db'),connect_args={'check_same_thread':False})
    Base.metadata.create_all(engine);factory=sessionmaker(bind=engine)
    with factory() as db:
        user=User(email='app-ui@example.test',name='UI Student');db.add(user);db.flush();user_id=user.id
        upload=Upload(user_id=user_id,filename='manual.docx',original_filename='manual.docx',file_path=str(manual),file_type='docx',file_size=manual.stat().st_size)
        db.add(upload);db.flush();upload_id=upload.id
        workflow=AssignmentWorkflow(upload_id=upload.id,status='processing',stage='generate',language='auto',
            questions=[{'id':1,'text':objective,'language':language}],results=[],output_name='app_ui_result')
        db.add(workflow);db.commit();workflow_id=workflow.id
    if language=='python':
        response={'answer':'Adds and validates on the Flask server.','entrypoint':'app.py','readme':FLASK_FILES['README.md'],
            'source_capture_files':['app.py','templates/index.html'],'files':[{'path':n,'content':v} for n,v in FLASK_FILES.items() if n!='README.md'],
            'browser_tests':[{k:v for k,v in t.items() if k!='version'} for t in FLASK_TESTS]}
    else:
        java=SWING_SOURCE.replace('Integer.parseInt(first.getText()) + Integer.parseInt(second.getText())','23') if wrong else SWING_SOURCE
        response={'answer':'Actual Swing event handlers.','entrypoint':'Calculator.java','readme':'Run javac Calculator.java then java Calculator.',
            'files':[{'path':'Calculator.java','content':java}],'source_capture_files':['Calculator.java'],
            'gui_tests':[{k:v for k,v in t.items() if k!='version'} for t in GUI_TESTS]}
    with patch('app.services.assignment_workflow.SessionLocal',factory),\
         patch('app.services.assignment_workflow.generation_service.generate_json',AsyncMock(side_effect=lambda *a,**k:deepcopy(response))),\
         patch('app.services.assignment_workflow.settings.REPORT_DIR',str(tmp_path/'reports')),\
         patch('app.services.assignment_workflow.settings.SCREENSHOT_DIR',str(tmp_path/'screenshots')),\
         patch('app.services.assignment_workflow.settings.REACT_TEMP_DIR',str(tmp_path/'runtime')):
        await process_assignment(workflow_id)
    with factory() as db:
        workflow=db.get(AssignmentWorkflow,workflow_id);assert workflow.status==('failed' if wrong else 'completed'),workflow.error
        from app.models import CreditWallet, WorkflowCreditHold
        balance=db.get(CreditWallet,(user_id,settings.RAZORPAY_MODE))
        assert balance.balance==(settings.BILLING_TRIAL_CREDITS if wrong else settings.BILLING_TRIAL_CREDITS-1)
        assert balance.reserved==0 and db.get(WorkflowCreditHold,workflow_id).state=='settled'
        assert workflow.report_id and workflow.results[0]['verification']['status']==('failed' if wrong else 'passed')
        report=db.get(Report,workflow.report_id);document=Document(report.file_path)
        text='\n'.join(p.text for p in document.paragraphs)
        assert 'VS Code-style source' in text
        if wrong:assert 'Diagnostic report' in text
        if language=='java':
            assert any(e['kind']=='gui' for e in workflow.results[0]['evidence_plan']['evidence']['selected'])
            assert 'Java UI interaction verification' in text
    from app.routers import workflows,download
    app=FastAPI();app.include_router(workflows.router,prefix='/api');app.include_router(download.router,prefix='/api')
    def database():
        with factory() as db:yield db
    app.dependency_overrides[get_db]=database
    app.dependency_overrides[get_current_user]=lambda:SimpleNamespace(id=user_id,email='app-ui@example.test',name='UI Student')
    with TestClient(app) as client:
        status=client.get(f'/api/workflows/{upload_id}');assert status.status_code==200
        downloaded=client.get(status.json()['report']['download_url']);assert downloaded.status_code==200
        with ZipFile(BytesIO(downloaded.content)) as package:
            media=[n for n in package.namelist() if n.startswith('word/media/')];assert len(media)>=3
            for name in media:
                with Image.open(BytesIO(package.read(name))) as image:image.verify()
        preview=client.get(f'/api/workflows/{upload_id}/preview');assert preview.status_code==200
        assert 'data:image/png;base64,' in preview.text
    engine.dispose()
