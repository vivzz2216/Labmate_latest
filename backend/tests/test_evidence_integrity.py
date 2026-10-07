"""Execution evidence must be genuine and a failed run must not become a report."""

import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from docx import Document
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import AssignmentWorkflow, Upload, User
from app.security.generated_code import validate_generated_code
from app.services.runtime_engine import runtime_engine
from app.services.assignment_workflow import process_assignment


class ExecutionEvidenceTests(unittest.TestCase):
    def test_python_pipe_echoes_actual_consumed_values_in_order(self):
        source = "a = int(input('First: '))\nb = int(input('Second: '))\nprint('Sum:', a+b)"
        with tempfile.TemporaryDirectory() as directory, patch("app.services.runtime_engine.settings.REACT_TEMP_DIR", directory):
            result = asyncio.run(runtime_engine.execute(source, "python", stdin="25\n5\n"))
        self.assertTrue(result.success, result.error)
        self.assertIn("First: 25\nSecond: 5\nSum: 30", result.output.replace("\r\n", "\n"))

    def test_super_initializer_is_allowed_but_introspection_is_not(self):
        validate_generated_code("class Child(Base):\n    def __init__(self):\n        super().__init__()", "python")
        with self.assertRaisesRegex(ValueError, "introspection"):
            validate_generated_code("print(object.__class__)", "python")

    def test_c_terminal_transcript_interleaves_stdin_values(self):
        from app.services.screenshot_service import screenshot_service
        code = '#include <stdio.h>\nint main() {\n  int a, b;\n  printf("Enter first number: ");\n  scanf("%d", &a);\n  printf("Enter second number: ");\n  scanf("%d", &b);\n  printf("Sum: %d\\n", a + b);\n  return 0;\n}'
        raw_output = "Enter first number: Enter second number: Sum: 42\r\n"
        stdin_data = "25\n17\n"
        result = screenshot_service._interleave_stdin_output(raw_output, stdin_data, source_code=code)
        self.assertEqual(result, "Enter first number: 25\nEnter second number: 17\nSum: 42")

    def test_cpp_terminal_transcript_interleaves_stdin_values(self):
        from app.services.screenshot_service import screenshot_service
        code = '#include <iostream>\nusing namespace std;\nint main() {\n  int a, b;\n  cout << "Enter first number: ";\n  cin >> a;\n  cout << "Enter second number: ";\n  cin >> b;\n  cout << "Sum: " << a + b << endl;\n  return 0;\n}'
        raw_output = "Enter first number: Enter second number: Sum: 42\r\n"
        stdin_data = "25\n17\n"
        result = screenshot_service._interleave_stdin_output(raw_output, stdin_data, source_code=code)
        self.assertEqual(result, "Enter first number: 25\nEnter second number: 17\nSum: 42")

    def test_java_terminal_transcript_interleaves_stdin_values(self):
        from app.services.screenshot_service import screenshot_service
        code = 'import java.util.Scanner;\npublic class Sum {\n  public static void main(String[] args) {\n    Scanner sc = new Scanner(System.in);\n    System.out.print("Enter first number: ");\n    int a = sc.nextInt();\n    System.out.print("Enter second number: ");\n    int b = sc.nextInt();\n    System.out.println("Sum: " + (a + b));\n  }\n}'
        raw_output = "Enter first number: Enter second number: Sum: 42\r\n"
        stdin_data = "25\n17\n"
        result = screenshot_service._interleave_stdin_output(raw_output, stdin_data, source_code=code)
        self.assertEqual(result, "Enter first number: 25\nEnter second number: 17\nSum: 42")

    def test_interleave_stdin_output_is_idempotent_on_repeated_calls(self):
        from app.services.screenshot_service import screenshot_service
        code = '#include <stdio.h>\nint main() {\n  int a, b;\n  printf("Enter first number: ");\n  scanf("%d", &a);\n  printf("Enter second number: ");\n  scanf("%d", &b);\n  printf("Sum: %d\\n", a + b);\n  return 0;\n}'
        raw_output = "Enter first number: Enter second number: Sum: 42\r\n"
        stdin_data = "25\n17\n"
        first_pass = screenshot_service._interleave_stdin_output(raw_output, stdin_data, source_code=code)
        self.assertEqual(first_pass, "Enter first number: 25\nEnter second number: 17\nSum: 42")
        second_pass = screenshot_service._interleave_stdin_output(first_pass, stdin_data, source_code=code)
        self.assertEqual(second_pass, first_pass)
        third_pass = screenshot_service._interleave_stdin_output(second_pass, stdin_data, source_code=code)
        self.assertEqual(third_pass, first_pass)

    def test_embed_screenshot_continuation_slices_have_titled_captions(self):
        from PIL import Image
        from app.services.docx_layout import embed_screenshot
        with tempfile.TemporaryDirectory() as folder:
            image_path = Path(folder) / "tall_code.png"
            # Create a 1000x2500 white image with rows of text
            img = Image.new("RGB", (1000, 2500), "white")
            for y in range(0, 2500, 50):
                for x in range(100, 800):
                    img.putpixel((x, y), (20, 20, 20))
            img.save(image_path, dpi=(96, 96))
            doc = Document()
            embed_screenshot(doc, image_path, caption="Program 1 — Code")
            # Should have multiple slices with captions for each slice
            captions = [p.text for p in doc.paragraphs if "Program 1" in p.text]
            self.assertGreater(len(captions), 1)
            self.assertIn("Part 1 of", captions[0])
            self.assertIn("continued", captions[1])


class FailedReportGateTests(unittest.IsolatedAsyncioTestCase):
    async def test_failed_execution_cannot_become_completed_word_report(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = Document()
            source.add_paragraph("Write a Python program to print a greeting.")
            source_path = root / "manual.docx"
            source.save(source_path)
            engine = create_engine("sqlite:///:memory:")
            Base.metadata.create_all(engine)
            factory = sessionmaker(bind=engine)
            with factory() as db:
                user = User(email="evidence@example.test", name="Evidence Student")
                db.add(user)
                db.flush()
                upload = Upload(user_id=user.id, filename="manual.docx", original_filename="manual.docx",
                                file_path=str(source_path), file_type="docx", file_size=source_path.stat().st_size)
                db.add(upload)
                db.flush()
                workflow = AssignmentWorkflow(upload_id=upload.id, status="processing", stage="generate",
                    questions=[{"id": 1, "text": "Write a Python program to print a greeting.", "language": "python"}],
                    results=[], language="auto", output_name="must_not_exist")
                db.add(workflow)
                db.commit()
                workflow_id = workflow.id
            generated = {"answer": "Print a greeting.", "code": "print('Hello')", "language": "python", "stdin": ""}
            with patch("app.services.assignment_workflow.SessionLocal", factory), \
                 patch("app.services.assignment_workflow.generation_service.generate_json", AsyncMock(return_value=generated)), \
                 patch("app.services.assignment_workflow.execute_solution", AsyncMock(return_value=(False, "", "Execution rejected", []))), \
                 patch("app.services.assignment_workflow.screenshot_service.generate_screenshot", AsyncMock(return_value=(False, "", 0, 0))), \
                 patch("app.services.assignment_workflow.settings.REPORT_DIR", str(root / "reports")):
                await process_assignment(workflow_id)
            with factory() as db:
                result = db.get(AssignmentWorkflow, workflow_id)
                self.assertEqual(result.status, "failed")
                self.assertIsNone(result.report_id)
                self.assertIn("verified report", result.error)
            engine.dispose()


class WorkflowInterleaveRegressionTests(unittest.IsolatedAsyncioTestCase):
    async def test_c_workflow_interleaves_inputs_exactly_once_without_duplicates(self):
        from app.services.screenshot_service import screenshot_service
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = Document()
            source.add_paragraph("Write a C program to add two numbers.")
            source_path = root / "manual.docx"
            source.save(source_path)
            engine = create_engine("sqlite:///:memory:")
            Base.metadata.create_all(engine)
            factory = sessionmaker(bind=engine)
            with factory() as db:
                user = User(email="c_student@example.test", name="C Student")
                db.add(user)
                db.flush()
                upload = Upload(user_id=user.id, filename="manual.docx", original_filename="manual.docx",
                                file_path=str(source_path), file_type="docx", file_size=source_path.stat().st_size)
                db.add(upload)
                db.flush()
                workflow = AssignmentWorkflow(upload_id=upload.id, status="processing", stage="generate",
                    questions=[{"id": 1, "text": "Write a C program to add two numbers.", "language": "c"}],
                    results=[], language="auto", output_name="c_lab_report")
                db.add(workflow)
                db.commit()
                workflow_id = workflow.id

            code = '#include <stdio.h>\nint main() {\n  int a, b;\n  printf("Enter first number: ");\n  scanf("%d", &a);\n  printf("Enter second number: ");\n  scanf("%d", &b);\n  printf("Sum: %d\\n", a + b);\n  return 0;\n}'
            raw_output = "Enter first number: Enter second number: Sum: 42\r\n"
            generated = {"answer": "Program adds two numbers.", "code": code, "language": "c", "stdin": "25\n17\n"}
            dummy_screenshot = str(root / "shot.png")
            Image.new("RGB", (100, 100), "white").save(dummy_screenshot)

            captured_calls = []
            async def mock_generate_screenshot(code_arg, output_arg, theme_arg, *args, **kwargs):
                captured_calls.append({
                    "code": code_arg,
                    "output": output_arg,
                    "theme": theme_arg,
                    "input_echoed": kwargs.get("input_echoed"),
                    "view_mode": kwargs.get("view_mode"),
                    "stdin_data": kwargs.get("stdin_data"),
                })
                # Call real _render_template to verify HTML does not contain duplicate 25 25
                html = await screenshot_service._render_template(
                    highlighted_code=code_arg,
                    output=output_arg,
                    theme=theme_arg,
                    source_code=code_arg,
                    stdin_data=kwargs.get("stdin_data"),
                    view_mode=kwargs.get("view_mode", "output"),
                    input_echoed=kwargs.get("input_echoed", False),
                )
                self.assertNotIn("25 25", html)
                self.assertNotIn("17 17", html)
                return True, dummy_screenshot, 1000, 600

            with patch("app.services.assignment_workflow.SessionLocal", factory), \
                 patch("app.services.assignment_workflow.generation_service.generate_json", AsyncMock(return_value=generated)), \
                 patch("app.services.assignment_workflow.execute_solution", AsyncMock(return_value=(True, raw_output, "", []))), \
                 patch("app.services.assignment_workflow.screenshot_service.generate_screenshot", side_effect=mock_generate_screenshot), \
                 patch("app.services.assignment_workflow.settings.REPORT_DIR", str(root / "reports")):
                await process_assignment(workflow_id)

            with factory() as db:
                result = db.get(AssignmentWorkflow, workflow_id)
                self.assertEqual(result.status, "completed")
                self.assertEqual(len(result.results), 1)
                res = result.results[0]
                # Check workflow output has single interleave
                self.assertIn("Enter first number: 25", res["output"])
                self.assertIn("Enter second number: 17", res["output"])
                self.assertNotIn("25 25", res["output"])
                self.assertNotIn("17 17", res["output"])
                self.assertNotIn("2525", res["output"])

            # Verify both editor and output views were captured
            view_modes = [c["view_mode"] for c in captured_calls]
            self.assertEqual(view_modes, ["editor", "output"])

            # Verify input_echoed was passed as True to prevent double interleaving
            self.assertTrue(all(c["input_echoed"] is True for c in captured_calls))

            # Verify that screenshot output argument did NOT contain duplicate 25 25
            for c in captured_calls:
                self.assertNotIn("25 25", c["output"])
                self.assertNotIn("17 17", c["output"])
            engine.dispose()
