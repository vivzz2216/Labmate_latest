import asyncio
import time
from unittest.mock import MagicMock, patch
import pytest

from app.monitoring import EventLoopLagMonitor, PerformanceMonitor, get_db_pool_status


@pytest.mark.asyncio
async def test_event_loop_lag_monitor_measurement():
    monitor = EventLoopLagMonitor(interval=0.02)
    await monitor.start()
    try:
        # Give it a moment to take samples
        await asyncio.sleep(0.08)
        metrics = monitor.get_metrics()
        assert metrics["total_samples"] >= 2
        assert "current_lag_ms" in metrics
        assert "peak_lag_ms" in metrics
        assert "avg_lag_ms" in metrics

        # Induce an intentional blocking delay on the event loop
        time.sleep(0.06)
        await asyncio.sleep(0.05)
        updated = monitor.get_metrics()
        # Peak lag should record at least ~40ms
        assert updated["peak_lag_ms"] >= 30.0
    finally:
        await monitor.stop()
        assert not monitor.running


def test_get_db_pool_status_sqlite():
    status = get_db_pool_status()
    assert isinstance(status, dict)
    assert "pool_type" in status


def test_get_db_pool_status_queue_pool():
    mock_pool = MagicMock()
    mock_pool.size.return_value = 5
    mock_pool.checkedin.return_value = 4
    mock_pool.checkedout.return_value = 1
    mock_pool.overflow.return_value = 0
    type(mock_pool).__name__ = "QueuePool"

    mock_engine = MagicMock()
    mock_engine.pool = mock_pool

    with patch("app.database.engine", mock_engine):
        status = get_db_pool_status()
        assert status["pool_type"] == "QueuePool"
        assert status["size"] == 5
        assert status["checkedin"] == 4
        assert status["checkedout"] == 1
        assert status["overflow"] == 0


@pytest.mark.asyncio
async def test_performance_monitor_integration():
    pm = PerformanceMonitor()
    pm.record_request(0.05, success=True)
    pm.record_request(1.5, success=False)

    await pm.start_monitoring(interval=0.02)
    try:
        await asyncio.sleep(0.06)
        metrics = pm.get_metrics()
        assert metrics["requests_total"] == 2
        assert metrics["requests_success"] == 1
        assert metrics["requests_error"] == 1
        assert "event_loop_lag" in metrics
        assert "db_pool" in metrics
    finally:
        await pm.stop_monitoring()


@pytest.mark.asyncio
async def test_workflow_preview_offload(tmp_path):
    from docx import Document
    from app.routers.workflows import _convert_docx_to_html

    doc_path = str(tmp_path / "sample.docx")
    doc = Document()
    doc.add_paragraph("Hello Event Loop Offload")
    doc.save(doc_path)

    # Calling via asyncio.to_thread should execute without blocking
    result = await asyncio.to_thread(_convert_docx_to_html, doc_path)
    assert "Hello Event Loop Offload" in result


@pytest.mark.asyncio
async def test_upload_preview_offload(tmp_path):
    from docx import Document
    from app.routers.upload import _render_docx_preview

    doc_path = str(tmp_path / "upload_sample.docx")
    doc = Document()
    doc.add_paragraph("Upload Preview Test")
    doc.save(doc_path)

    result = await asyncio.to_thread(_render_docx_preview, doc_path)
    assert "Upload Preview Test" in result
