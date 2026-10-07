import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from fastapi import BackgroundTasks, Request
from PIL import Image
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.database import Base
from app.models import AssignmentWorkflow, Report, Upload, User
from app.routers.workflows import ProcessInput, process
from app.services.assignment_workflow import process_assignment
from app.services.generation_service import GenerationUnavailableError


class WorkflowResumeTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.folder = tempfile.TemporaryDirectory()
        previous_capture = Path(self.folder.name) / "previous_capture.png"
        Image.new("RGB", (120, 60), "white").save(previous_capture)
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.factory = sessionmaker(bind=self.engine)
        self.questions = [
            {"id": 1, "text": "Write a Python program to calculate factorial of 5.", "language": "python"},
            {"id": 2, "text": "Write a Python program to reverse a string.", "language": "python"},
        ]
        with self.factory() as db:
            user = User(email="resume@example.com", name="Resume Test")
            db.add(user); db.flush()
            upload = Upload(user_id=user.id, filename="manual.docx", original_filename="manual.docx", file_path="unused", file_type="docx", file_size=1)
            db.add(upload); db.flush()
            workflow = AssignmentWorkflow(upload_id=upload.id, status="processing", stage="generate", questions=self.questions,
                results=[{"id": 1, "status": "completed", "output": "120", "answer": "Factorial", "code": "print(120)", "question": self.questions[0]["text"], "language": "python", "error": "", "screenshot_paths": [str(previous_capture)]}],
                language="auto", instructions="", output_name="report", progress=52)
            db.add(workflow); db.commit()
            self.id, self.upload_id, self.user_id = workflow.id, upload.id, user.id

    async def asyncTearDown(self):
        self.engine.dispose()
        self.folder.cleanup()

    async def test_outage_saves_checkpoint_and_resume_keeps_completed_results(self):
        with patch("app.services.assignment_workflow.SessionLocal", self.factory), patch("app.services.assignment_workflow.document_text", return_value="Programming lab"), patch("app.services.assignment_workflow.generation_service.generate_json", AsyncMock(side_effect=GenerationUnavailableError("Temporary outage"))) as generation:
            await process_assignment(self.id)
        self.assertEqual(generation.await_count, 1)  # question 1 is not regenerated
        with self.factory() as db:
            workflow = db.get(AssignmentWorkflow, self.id)
            self.assertEqual(workflow.status, "paused")
            self.assertEqual(workflow.results[0]["output"], "120")
            self.assertEqual(len(workflow.results), 1)
            payload = ProcessInput(questions=self.questions, language="auto", output_name="report")
            user = db.get(User, self.user_id)
            response = await process(self.upload_id, payload, Request({"type": "http", "method": "POST"}), BackgroundTasks(), user, db, True)
            self.assertEqual(response["status"], "queued")
            self.assertEqual(len(response["results"]), 1)
            self.assertEqual(response["progress"], 52)

    async def test_resume_executes_remaining_question_and_builds_word(self):
        from docx import Document
        root = Path(self.folder.name)
        source = Document()
        source.add_paragraph("Original assignment retained")
        source_path = root / "source.docx"
        source.save(source_path)
        with self.factory() as db:
            upload = db.get(Upload, self.upload_id)
            upload.file_path = str(source_path)
            db.commit()
        answer = {"answer": "Slice the string with a negative step.", "code": "print('hello'[::-1])", "language": "python"}
        with patch("app.services.assignment_workflow.SessionLocal", self.factory), patch("app.services.assignment_workflow.generation_service.generate_json", AsyncMock(return_value=answer)) as generation, patch("app.services.assignment_workflow.settings.REPORT_DIR", str(root / "reports")), patch("app.services.assignment_workflow.settings.SCREENSHOT_DIR", str(root / "screenshots")), patch("app.services.assignment_workflow.settings.REACT_TEMP_DIR", str(root / "runtime")):
            await process_assignment(self.id)
        with self.factory() as db:
            workflow = db.get(AssignmentWorkflow, self.id)
            self.assertEqual(workflow.status, "completed", workflow.error)
            self.assertEqual(generation.await_count, 1)
            self.assertEqual(workflow.results[1]["output"].strip(), "olleh")
            report = db.get(Report, workflow.report_id)
            final = Document(report.file_path)
            self.assertEqual(final.paragraphs[0].text, "Original assignment retained")
            self.assertEqual(len(final.inline_shapes), 3)  # prior capture, then new IDLE editor and shell
