from types import SimpleNamespace

from docx import Document

from app.services.assignment_workflow import document_text
from app.services.manual_sections import numbered_experiments


def test_adjacent_table_cells_preserve_heading_and_requirements(tmp_path):
    doc = Document()
    doc.add_paragraph('Experiment 1: Python counting')
    doc.add_paragraph('Aim: Write a Python program counting words.')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Experiment 24: React counter'
    table.cell(0, 1).text = 'Aim: Create a React counter. Two clicks must display 2.'
    doc.add_paragraph('Experiment 30: Node.js API')
    doc.add_paragraph('Aim: Build an Express API returning status 200.')
    path = tmp_path / 'manual.docx'
    doc.save(path)
    text = document_text(SimpleNamespace(file_type='docx', file_path=str(path)))
    tasks = numbered_experiments(text)
    assert [(q['id'], q['language']) for q in tasks] == [(1,'python'),(24,'react'),(30,'node')]
    assert 'Two clicks must display 2.' in tasks[1]['text']
    assert 'Experiment 30' not in tasks[1]['text']


def test_merged_experiment_heading_is_not_duplicated(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = 'Experiment 21: HTML task board'
    table.cell(1, 0).text = 'Aim:'
    table.cell(1, 1).text = 'Build an HTML task board with editing and search.'
    path = tmp_path / 'merged.docx'
    doc.save(path)
    text = document_text(SimpleNamespace(file_type='docx', file_path=str(path)))
    tasks = numbered_experiments(text)
    assert len(tasks) == 1 and tasks[0]['id'] == 21
    assert tasks[0]['language'] == 'html'
    assert 'Build an HTML task board with editing and search.' in tasks[0]['text']
