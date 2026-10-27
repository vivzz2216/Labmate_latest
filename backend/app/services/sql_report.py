"""Submission reports derived from MySQL receipts, without copied manual defects."""
import os
import re
import uuid
from pathlib import Path
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt, RGBColor
from ..config import settings
from ..security.validators import normalize_output_name
from .docx_layout import apply_student_metadata, save_document_atomic


def report_number(filename, question):
    m = re.search(r'exp[_ -]*(\d+)', filename, re.I) or re.search(r'experiment\s*[-–—]?\s*(\d+)', question, re.I)
    return int(m.group(1)) if m else None


def label_sql_comments(code, number):
    if number is None:return code
    return re.sub(r'^(\s*(?:--|\*)[^\n]*\bExperiment\s*[:–—-]?\s*)\d+\b',
                  lambda m:m.group(1)+str(number), code, flags=re.M|re.I)


def post_questions(text):
    start = re.search(r'\bPost[- ]+Experiment(?:al|s)?\s+Exercise[^\n]*', text, re.I)
    if not start:return []
    body = re.split(r'\b(?:Conclusion|References)\s*:', text[start.end():], maxsplit=1, flags=re.I)[0]
    body = re.sub(r'(?<=[.?])\s*\d+[.)]\s*', '\n', body)
    result = []
    for line in body.splitlines():
        line = re.sub(r'^\s*(?:[A-Z][.)]\s*)?(?:\d+[.)]\s*)?', '', line).strip()
        if re.match(r'^(?:what|why|how|explain|write syntax|difference|perform)\b', line, re.I):result.append(line)
    return result


def answer_question(question):
    q = question.lower()
    if 'drop' in q and 'truncate' in q and 'delete' in q:
        return 'DELETE removes selected rows using WHERE and can be rolled back in an uncommitted InnoDB transaction. TRUNCATE removes every row, retains the structure and resets auto-increment. DROP removes the definition and its data. MySQL DROP and TRUNCATE implicitly commit.'
    if 'drop' in q and 'truncate' in q:
        return 'DROP removes a table and its definition. TRUNCATE empties the table while retaining its structure and resets auto-increment. Both cause implicit commits in MySQL.'
    if 'rollback' in q and 'ddl' in q:
        return 'Most MySQL DDL statements cause an implicit commit, so a later ROLLBACK cannot undo them. Atomic DDL provides crash consistency, not user-controlled rollback of DDL.'
    if 'execution pattern' in q:
        return 'Logical query order is FROM and JOIN, WHERE, GROUP BY, HAVING, SELECT, DISTINCT, ORDER BY, then LIMIT. The optimizer may choose a different physical execution plan that preserves the result.'
    if 'ddl' in q and 'dml' in q:
        return 'DDL defines objects using CREATE, ALTER, RENAME, TRUNCATE and DROP. DML reads or changes records using SELECT, INSERT, UPDATE and DELETE. Transaction behavior depends on the statement and storage engine.'
    if 'access control' in q and 'role based' not in q:
        return 'Access control determines which accounts may access objects and which operations they may perform. GRANT assigns privileges and REVOKE removes them. Object and column privileges support least-privilege access.'
    if 'role based' in q:
        return 'Role-based access control groups privileges into roles assigned to accounts. A MySQL account activates an assigned role to use its privileges.'
    if 'tcl' in q:
        return 'START TRANSACTION begins an InnoDB transaction, COMMIT makes changes permanent, ROLLBACK discards uncommitted changes, and SAVEPOINT permits partial rollback. For example, insert a record inside a transaction, inspect it with SELECT, then roll back and inspect again. Keep DDL outside rollback demonstrations because it implicitly commits.'
    if 'avoid' in q and 'trigger' in q:
        return 'Avoid triggers for behavior better expressed by constraints, external side effects and complex hidden workflows. Row triggers add write latency and complicate debugging. Avoid recursive dependencies and updates to the table already being changed.'
    if 'call' in q and 'procedure' in q:
        return "A no-parameter procedure is invoked using CALL followed by its name and empty parentheses. A procedure with an IN parameter needs a compatible argument, such as CALL exam('JEE') when that value exists in the example data."
    if ('delete' in q or 'alter' in q) and 'procedure' in q:
        return 'DROP PROCEDURE IF EXISTS procedure_name removes a procedure. ALTER PROCEDURE changes characteristics such as its comment or SQL SECURITY setting. Changing parameters or the body requires dropping and recreating the procedure in MySQL.'
    return None


def clean_report(upload, workflow, source):
    doc = Document();section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11.7), Inches(8.3)
    section.top_margin = section.bottom_margin = Inches(.5)
    section.left_margin = section.right_margin = Inches(.6)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        doc.styles[name].font.name = 'Calibri';doc.styles[name].font.color.rgb = RGBColor(0,0,0)
        for border in doc.styles[name]._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Heading 1'].font.size = Pt(14)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)
    for ri,result in enumerate(workflow.results):
        number = report_number(upload.original_filename if len(workflow.results)==1 else '',result.get('question','')) or result['id']
        heading = result.get('question','').split('\n')[0]
        title = re.sub(r'^\s*Experiment\s*[-–—]?\s*\d+[\s.:–—-]*','',heading,flags=re.I).strip()
        title_paragraph=doc.add_heading(f'Experiment {number} {title}',0)
        if ri:title_paragraph.paragraph_format.page_break_before=True
        for label,key in [('Student','student_name'),('Roll number','roll_no'),('Department','department'),('Institution','institution')]:
            value = getattr(workflow,key,'')
            if value and value not in {'Student','owner','admin','USN-2024-001','Engineering College','Computer Science & Engineering'}:doc.add_paragraph(f'{label}: {value}')
        doc.add_heading('Aim',1)
        aim = re.search(r'^(?:\d+[.)]\s*)?Aim\s*:\s*(.+)',result.get('question',''),re.M|re.I)
        doc.add_paragraph(aim.group(1).strip() if aim else title.rstrip('.'))
        v = result.get('verification') or {};steps = v.get('statements') or []
        ok = result.get('status')=='completed' and v.get('verified') is True
        queries = [s for s in steps if s.get('success') and s.get('columns')]
        doc.add_paragraph(f"Database environment: {v.get('engine','MySQL')}.")
        doc.add_heading('Result and Observation',1)
        doc.add_paragraph(f'{len(steps)} statements executed; {len(queries)} query outputs recorded. '+('The saved execution and required result checks succeeded.' if ok else 'Diagnostic report: the experiment is incomplete; the details require correction.'))
        for check in v.get('checks',[]):
            if not check.get('passed'):doc.add_paragraph('Correction required: '+check.get('detail',''))
        if not ok and result.get('error'):doc.add_paragraph(str(result['error'])[:2000])
        table = doc.add_table(rows=1,cols=3);table.style='Light Shading Accent 1'
        for cell,text in zip(table.rows[0].cells,('Query','Executed statement','Returned rows')):cell.text=text
        for index,step in enumerate(queries,1):
            row=table.add_row();sql=' '.join(step['statement'].split())
            for cell,text in zip(row.cells,(str(index),sql if len(sql)<150 else sql[:147]+'...',str(len(step.get('rows',[]))))):cell.text=text
        table.autofit=False
        widths=(.55,8.95,.9)
        for column,width in zip(table.columns,widths):column.width=Inches(width)
        for row in table.rows:
            for cell,width in zip(row.cells,widths):
                cell.width=Inches(width)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before=Pt(0)
                    paragraph.paragraph_format.space_after=Pt(2)
                    paragraph.paragraph_format.line_spacing=1.0
                    for run in paragraph.runs:run.font.size=Pt(9)
        questions=post_questions(source)
        if questions:
            questions_heading=doc.add_heading('Post Experiment Questions',1)
            questions_heading.paragraph_format.page_break_before=True
            for index,q in enumerate(questions,1):
                answer=answer_question(q)
                if not answer:raise ValueError('A completed SQL report needs an answer to: '+q)
                doc.add_paragraph(f'{index}. {q}',style='Heading 2');doc.add_paragraph(answer)
        doc.add_heading('Conclusion',1)
        doc.add_paragraph((f'The experiment demonstrated {title.rstrip(".").lower()} on MySQL. The recorded results contain {len(queries)} query outputs and the required execution checks succeeded.') if ok else 'The experiment did not meet every required execution or result check. The diagnostic evidence must be corrected before submission.')
        frames=v.get('evidence_frames') or [{'path':p,'kind':'database','caption':c} for p,c in zip(result.get('screenshot_paths',[]),result.get('screenshot_captions',[]))]
        for frame in frames:
            if not os.path.isfile(frame['path']):raise ValueError('Required SQL evidence image is missing: '+frame['path'])
            evidence_heading=doc.add_heading('SQL Program' if frame.get('kind')=='source_preview' else 'Output',1)
            evidence_heading.paragraph_format.page_break_before=True
            caption=re.sub(r'\s*\[.*?Workbench-style rendered interface\]','',frame.get('caption','MySQL execution evidence'))
            caption=caption.replace('MySQL Workbench SQL Script',f'Experiment {number} SQL script').replace('MySQL Workbench Query','Query').replace('MySQL Workbench Execution Step','Execution step')
            paragraph=doc.add_paragraph(caption)
            paragraph.paragraph_format.keep_with_next=True
            for run in paragraph.runs:run.font.size=Pt(9)
            doc.add_picture(frame['path'],width=Inches(9.8))
            image_paragraph=doc.paragraphs[-1]
            image_paragraph.paragraph_format.line_spacing=1.0
            image_paragraph.paragraph_format.space_before=Pt(0)
            image_paragraph.paragraph_format.space_after=Pt(0)
    name=normalize_output_name(workflow.output_name or Path(upload.original_filename).stem,max_length=80)+'.docx'
    directory=Path(settings.REPORT_DIR)/'workflows'/str(workflow.id);directory.mkdir(parents=True,exist_ok=True)
    path=directory/(uuid.uuid4().hex[:8]+'_'+name)
    apply_student_metadata(doc,getattr(workflow,'student_name',''),title='SQL Laboratory Report')
    save_document_atomic(doc,str(path))
    from .storage_service import storage_service
    storage_service.sync_file(str(path))
    return str(path),name
