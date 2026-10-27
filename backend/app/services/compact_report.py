"""Short report layout using saved evidence, with no model calls."""
import re
import uuid
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches,Pt,RGBColor

from ..config import settings
from ..security.validators import normalize_output_name
from .docx_layout import apply_student_metadata,embed_screenshot,save_document_atomic


def short_text(text,limit=180):
    text=re.sub(r'[*`#]+','',str(text or ''))
    text=' '.join(text.split())
    return text if len(text)<=limit else text[:limit-1].rsplit(' ',1)[0]+'…'


def diagnostic_text(result):
    lines=['Diagnostic report — this experiment has not passed verification. '+short_text(result.get('error') or 'Required execution evidence is missing.',300)]
    verification=result.get('verification') or {}
    kind=verification.get('kind','browser')
    if verification.get('status')=='failed':
        label={'notebook':'Python notebook verification','database':'MySQL verification',
            'gui':'Java UI interaction verification','http':'HTTP request verification'}.get(kind,'Browser interaction verification')
        lines.append(label+': FAILED')
        for test in verification.get('tests',[]):
            if test.get('status')!='failed':continue
            for index,step in enumerate((test.get('steps') or test.get('actions') or [])[:8],1):
                target=' '.join(str(step.get(key,'')) for key in ('locator_type','role','locator')).strip()
                lines.append(short_text(f"Step {index}: {step.get('action','')} {target}"))
            if test.get('expected_text'):lines.append('Expected visible text: '+short_text('; '.join(test['expected_text'])))
            if test.get('observed_text'):lines.append('Observed page text: '+short_text(test['observed_text']))
    if not result.get('screenshot_paths'):
        lines.append('Browser evidence unavailable.' if kind in {'browser','render','fullstack'} else 'Execution evidence unavailable.')
    return '\n'.join(lines)


def build_compact_report(upload,workflow):
    document=Document()
    section=document.sections[0]
    landscape=all(r.get('language')=='sql' for r in workflow.results)
    if landscape:section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width=Inches(11.7 if landscape else 8.5)
    section.page_height=Inches(8.3 if landscape else 11)
    section.left_margin=section.right_margin=Inches(.7)
    section.top_margin=section.bottom_margin=Inches(.6)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style=document.styles[name]
        style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
        for border in style._element.xpath('.//w:pBdr'):border.getparent().remove(border)
    document.styles['Normal'].font.size=Pt(11)
    document.styles['Heading 2'].font.size=Pt(13)
    for index,result in enumerate(workflow.results):
        question=next((line.strip() for line in result.get('question','').splitlines() if line.strip()),'')
        heading=document.add_paragraph(short_text(question or f"Program {result['id']}",240),style='Heading 2')
        heading.paragraph_format.keep_with_next=True
        heading.paragraph_format.page_break_before=index>0
        if result.get('status')!='completed':
            document.add_paragraph(diagnostic_text(result))
        paths=result.get('screenshot_paths') or []
        if not paths and result.get('status')=='completed':
            raise ValueError('Cannot export a completed exercise without its required screenshots.')
        captions=result.get('screenshot_captions') or []
        for image_index,path in enumerate(paths):
            if not Path(path).is_file():raise ValueError('A required screenshot is missing: '+str(path))
            caption=short_text(captions[image_index] if image_index<len(captions) else 'Execution evidence.') or 'Execution evidence.'
            embed_screenshot(document,path,caption=caption)
    name=normalize_output_name(workflow.output_name or 'lab_report',max_length=80)+'.docx'
    directory=Path(settings.REPORT_DIR)/'workflows'/str(workflow.id)
    directory.mkdir(parents=True,exist_ok=True)
    path=directory/(uuid.uuid4().hex[:8]+'_'+name)
    apply_student_metadata(document,getattr(workflow,'student_name',''),title='Laboratory Results')
    save_document_atomic(document,str(path))
    from .storage_service import storage_service
    storage_service.sync_file(str(path))
    return str(path),name
