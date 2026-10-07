"""
generate_final_boss_manuals.py
==============================
Generates 20 complete, realistic College Laboratory Manual Word documents (.docx):
 - 5 Python Manuals
 - 5 Java Manuals
 - 2 C Manuals
 - 3 C++ Manuals (DSA Focus: Trees, Graphs with user input, Stacks, Queues)
 - 5 Web Dev Manuals (HTML/CSS, JS DOM, React Components, React Hooks/State, Node.js REST & Cookies)

Each document contains:
 - Institution Header & Course Information
 - Section A: Theory & Viva Questions (to verify Labmate skips theory)
 - Section B: Laboratory Programming Exercises (actionable questions for Labmate to extract)
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_BASE = r"c:\Users\pilla\OneDrive\Desktop\labmate_latest\final_boss_manuals"

MANUALS_DATA = [
    # ══════════════════════════════════════════════════════════════════════════
    # 1. PYTHON (5 Manuals)
    # ══════════════════════════════════════════════════════════════════════════
    {
        "folder": "python",
        "filename": "python_manual_1_basics_control_flow.docx",
        "course": "CS201: Python Programming Laboratory",
        "title": "Lab Manual 01 — Basic Syntax, Data Types & Control Flow",
        "dept": "Department of Computer Science & Engineering",
        "theory_questions": [
            "Explain the difference between mutable and immutable data types in Python with examples.",
            "What is PEP 8 and why is it important in Python programming?",
            "Differentiate between break, continue, and pass statements in Python loops.",
            "Explain Python's dynamic typing and how type casting works."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Prime Number Checker and Interval Generator",
                "text": "Write a Python program to check whether an integer entered by the user is prime, and print all prime numbers in a user-specified interval [lower, upper] using for loops."
            },
            {
                "num": 2,
                "title": "Quadratic Equation Solver",
                "text": "Write a Python program to find and display all real and complex roots of a quadratic equation ax^2 + bx + c = 0 based on the discriminant using conditional branching."
            },
            {
                "num": 3,
                "title": "Number Pattern Printing",
                "text": "Write a Python program to print a Floyd's triangle and an inverted pyramid of numbers of height N entered by the user."
            }
        ]
    },
    {
        "folder": "python",
        "filename": "python_manual_2_functions_recursion.docx",
        "course": "CS201: Python Programming Laboratory",
        "title": "Lab Manual 02 — Modular Programming, Functions & Recursion",
        "dept": "Department of Computer Science & Engineering",
        "theory_questions": [
            "What are default arguments, keyword arguments, and arbitrary arguments (*args, **kwargs) in Python?",
            "Explain the scope resolution of variables in Python using the LEGB rule.",
            "What are anonymous or lambda functions? Mention their use cases with map() and filter().",
            "Explain the base condition in recursion and what causes a RecursionError."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Tower of Hanoi Recursive Solver",
                "text": "Write a Python program to solve the Tower of Hanoi puzzle for N disks using a recursive function and display every disc movement step."
            },
            {
                "num": 2,
                "title": "GCD and LCM via Euclidean Algorithm",
                "text": "Write a Python program with functions to compute the Greatest Common Divisor (GCD) and Least Common Multiple (LCM) of two numbers using Euclid's recursive algorithm."
            },
            {
                "num": 3,
                "title": "String Palindrome and Anagram Verifier",
                "text": "Write a Python program containing functions to test whether a given string is a palindrome and whether two user-supplied strings are anagrams of each other."
            }
        ]
    },
    {
        "folder": "python",
        "filename": "python_manual_3_data_structures.docx",
        "course": "CS201: Python Programming Laboratory",
        "title": "Lab Manual 03 — Built-in Data Structures & Comprehensions",
        "dept": "Department of Computer Science & Engineering",
        "theory_questions": [
            "Compare Lists, Tuples, Sets, and Dictionaries in terms of order, indexing, mutability, and duplicates.",
            "Explain list and dictionary comprehensions and compare them with standard loops in terms of readability and performance.",
            "What is shallow copying vs deep copying in Python? How does the copy module handle nested structures?",
            "Explain the working of zip() and enumerate() built-in functions with examples."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Word Frequency Counter with Dictionaries",
                "text": "Write a Python program to read a paragraph of text from the user, strip punctuation, convert to lowercase, and display the frequency count of each unique word sorted in descending order."
            },
            {
                "num": 2,
                "title": "Matrix Multiplication using List Comprehensions",
                "text": "Write a Python program to perform matrix multiplication of two user-defined matrices of size M x K and K x N and display the resultant matrix."
            },
            {
                "num": 3,
                "title": "Set Operations Simulator",
                "text": "Write a Python program to take two integer lists from the user, convert them into sets, and display their Union, Intersection, Difference (A-B), and Symmetric Difference."
            }
        ]
    },
    {
        "folder": "python",
        "filename": "python_manual_4_oop_exceptions.docx",
        "course": "CS201: Python Programming Laboratory",
        "title": "Lab Manual 04 — Object-Oriented Programming & Exception Handling",
        "dept": "Department of Computer Science & Engineering",
        "theory_questions": [
            "Explain the four pillars of OOP (Encapsulation, Abstraction, Inheritance, Polymorphism) in Python.",
            "What is Method Resolution Order (MRO) in Python multiple inheritance and how is it checked?",
            "Explain the try-except-else-finally blocks and their execution order during an exception.",
            "What are magic or dunder methods (__init__, __str__, __repr__, __add__) in Python classes?"
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Bank Account Management System",
                "text": "Write a Python program to create a BankAccount class with deposit, withdraw, and check_balance methods, implementing a custom InsufficientFundsException when withdrawal exceeds balance."
            },
            {
                "num": 2,
                "title": "Employee Hierarchy with Inheritance",
                "text": "Write a Python program to implement an Employee base class and derived classes Manager and Developer, overriding calculate_salary() to compute bonuses based on role."
            },
            {
                "num": 3,
                "title": "Complex Number Arithmetic via Operator Overloading",
                "text": "Write a Python program to create a Complex class and overload the + and * operators using __add__ and __mul__ dunder methods to add and multiply two complex numbers."
            }
        ]
    },
    {
        "folder": "python",
        "filename": "python_manual_5_file_handling_modules.docx",
        "course": "CS201: Python Programming Laboratory",
        "title": "Lab Manual 05 — File I/O, Serialization & Standard Modules",
        "dept": "Department of Computer Science & Engineering",
        "theory_questions": [
            "What is the difference between text mode and binary mode in Python file handling?",
            "Explain the advantages of using the with statement (context manager) when opening files.",
            "What is JSON serialization and deserialization? How do json.dump() and json.load() function?",
            "Explain the use of the os and sys modules in interacting with the host operating system."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Text File Line, Word, and Character Counter",
                "text": "Write a Python program to open a text file, count the total number of lines, words, uppercase letters, lowercase letters, and digits, and display the summary."
            },
            {
                "num": 2,
                "title": "Student CSV Record Processor",
                "text": "Write a Python program to create and write student records (RollNo, Name, Marks) to a CSV file, and then read the file to compute and display the highest scorer and average marks."
            },
            {
                "num": 3,
                "title": "JSON Configuration Loader and Validator",
                "text": "Write a Python program to read a JSON settings file, validate required fields using dictionary lookups, and handle FileNotFoundError and json.JSONDecodeError gracefully."
            }
        ]
    },

    # ══════════════════════════════════════════════════════════════════════════
    # 2. JAVA (5 Manuals)
    # ══════════════════════════════════════════════════════════════════════════
    {
        "folder": "java",
        "filename": "java_manual_1_fundamentals_arrays.docx",
        "course": "CS202: Object-Oriented Programming with Java",
        "title": "Lab Manual 01 — Java Fundamentals, Control Structures & Arrays",
        "dept": "Department of Information Technology",
        "theory_questions": [
            "Explain the role of JVM, JRE, and JDK in the Java execution environment.",
            "Why is Java considered platform independent and what is bytecode?",
            "Differentiate between primitive types and wrapper classes in Java.",
            "Explain the purpose of public static void main(String[] args) method signature in Java."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Matrix Transpose and Multiplication",
                "text": "Write a Java program to read two 2D integer arrays from standard input, perform matrix multiplication, compute the transpose of the resulting matrix, and display the output in matrix format."
            },
            {
                "num": 2,
                "title": "Command-Line Palindrome and Sieve of Eratosthenes",
                "text": "Write a Java program to accept an integer N from command-line arguments and print all prime numbers up to N using the Sieve of Eratosthenes algorithm."
            },
            {
                "num": 3,
                "title": "Array Sorting and Binary Search",
                "text": "Write a Java program to sort an array of student names in alphabetical order and implement binary search to locate a given query name."
            }
        ]
    },
    {
        "folder": "java",
        "filename": "java_manual_2_oop_inheritance.docx",
        "course": "CS202: Object-Oriented Programming with Java",
        "title": "Lab Manual 02 — Classes, Inheritance, Interfaces & Polymorphism",
        "dept": "Department of Information Technology",
        "theory_questions": [
            "Explain the difference between method overloading (static polymorphism) and method overriding (dynamic polymorphism).",
            "What is an abstract class and how does it differ from a Java Interface?",
            "Explain the super and this keywords in Java with constructor chaining examples.",
            "Why doesn't Java support multiple inheritance with classes, and how do default interface methods solve this?"
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Geometric Shape Hierarchy",
                "text": "Write a Java program to define an abstract class Shape with abstract methods area() and perimeter(), and implement subclasses Circle, Rectangle, and Triangle calculating and printing their respective metrics."
            },
            {
                "num": 2,
                "title": "Multiple Interface Implementation for Payables",
                "text": "Write a Java program demonstrating multiple interface inheritance where an Invoice class implements both Payable and Auditable interfaces with distinct method implementations."
            },
            {
                "num": 3,
                "title": "Vehicle Rental System Polymorphism",
                "text": "Write a Java program to create a Vehicle base class and derived classes Car, Bike, and Truck. Implement dynamic method dispatch to calculate rental charges based on distance and vehicle type."
            }
        ]
    },
    {
        "folder": "java",
        "filename": "java_manual_3_exception_multithreading.docx",
        "course": "CS202: Object-Oriented Programming with Java",
        "title": "Lab Manual 03 — Exception Handling & Multithreading",
        "dept": "Department of Information Technology",
        "theory_questions": [
            "Explain the hierarchy of Throwable, Error, Exception, RuntimeException in Java.",
            "Differentiate between checked and unchecked exceptions with code examples.",
            "What is the life cycle of a Thread in Java? Name all states.",
            "Explain thread synchronization and why the synchronized keyword or Lock interface is needed to prevent race conditions."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Custom Banking Exception Handler",
                "text": "Write a Java program to simulate a bank account with deposit and withdraw operations, throwing and handling a custom checked exception NegativeBalanceException when balance falls below zero."
            },
            {
                "num": 2,
                "title": "Producer-Consumer Problem using Multithreading",
                "text": "Write a Java program to solve the Producer-Consumer problem using inter-thread communication methods wait() and notify() with a shared synchronized buffer queue."
            },
            {
                "num": 3,
                "title": "Multi-threaded Odd-Even Number Printer",
                "text": "Write a Java program to create two threads (one printing odd numbers, one printing even numbers up to 20 in synchronized sequence) using the Runnable interface."
            }
        ]
    },
    {
        "folder": "java",
        "filename": "java_manual_4_collections_framework.docx",
        "course": "CS202: Object-Oriented Programming with Java",
        "title": "Lab Manual 04 — Java Collections Framework",
        "dept": "Department of Information Technology",
        "theory_questions": [
            "Explain the core interfaces of the Java Collections Framework (Collection, List, Set, Map, Queue).",
            "Differentiate between ArrayList and LinkedList in terms of memory layout and algorithmic complexity.",
            "How does HashMap work internally in Java? Explain hashing, buckets, and collision resolution.",
            "Differentiate between Comparable and Comparator interfaces in Java sorting."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Student Management System with ArrayList",
                "text": "Write a Java program using ArrayList<Student> to store student records (ID, Name, CGPA), allowing the user to insert, search, remove, and sort students by CGPA using a custom Comparator."
            },
            {
                "num": 2,
                "title": "Phone Book Directory with HashMap",
                "text": "Write a Java program to create a telephone directory using HashMap<String, String> that provides menu-driven options to add contacts, search phone numbers by name, and list all entries."
            },
            {
                "num": 3,
                "title": "Word Frequency Counter with TreeMap",
                "text": "Write a Java program to take a user sentence, split words, and count occurrences of each word using a TreeMap to display unique words in alphabetical order."
            }
        ]
    },
    {
        "folder": "java",
        "filename": "java_manual_5_file_io_streams.docx",
        "course": "CS202: Object-Oriented Programming with Java",
        "title": "Lab Manual 05 — File I/O, Character Streams & Object Serialization",
        "dept": "Department of Information Technology",
        "theory_questions": [
            "Explain byte streams (InputStream/OutputStream) vs character streams (Reader/Writer) in Java.",
            "What is Object Serialization in Java? What is the role of the Serializable interface and serialVersionUID?",
            "What does the transient keyword do during serialization?",
            "Explain try-with-resources statement and how AutoCloseable works."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "File Copy and Word/Char Statistics",
                "text": "Write a Java program to copy the contents of one text file into another destination file using BufferedReader and BufferedWriter, counting and printing total characters, words, and lines copied."
            },
            {
                "num": 2,
                "title": "Object Serialization & Deserialization",
                "text": "Write a Java program to serialize a Student object to a binary file student.ser using ObjectOutputStream and then deserialize it back to display student details."
            },
            {
                "num": 3,
                "title": "Log File Parser and Keyword Filter",
                "text": "Write a Java program to read an application log file line-by-line, extract and display only lines that contain the keyword 'ERROR' or 'WARN', writing them to error_log.txt."
            }
        ]
    },

    # ══════════════════════════════════════════════════════════════════════════
    # 3. C (2 Manuals)
    # ══════════════════════════════════════════════════════════════════════════
    {
        "folder": "c_cpp",
        "filename": "c_manual_1_arrays_strings_pointers.docx",
        "course": "CS101: Computer Programming in C",
        "title": "Lab Manual 01 — Pointers, Arrays, Strings & Dynamic Allocation",
        "dept": "Department of Computer Engineering",
        "theory_questions": [
            "Explain pointer arithmetic in C and how pointers relate to 1D arrays.",
            "Differentiate between call by value and call by reference in C.",
            "What is the difference between malloc(), calloc(), realloc(), and free() in dynamic memory management?",
            "Explain memory leak and dangling pointer issues in C with code snippets."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Array Reversal and Element Search using Pointers",
                "text": "Write a C program to reverse an integer array of N elements in-place and find a target element using pure pointer arithmetic without array index notation."
            },
            {
                "num": 2,
                "title": "String Manipulation without Built-in Library Functions",
                "text": "Write a C program to implement custom functions for string length (my_strlen), string concatenation (my_strcat), and string comparison (my_strcmp) using pointers."
            },
            {
                "num": 3,
                "title": "Dynamic 1D Array Allocation and Statistics",
                "text": "Write a C program to dynamically allocate memory for N floating-point numbers using malloc(), compute their mean, variance, and standard deviation, and free the allocated memory."
            }
        ]
    },
    {
        "folder": "c_cpp",
        "filename": "c_manual_2_structures_file_io.docx",
        "course": "CS101: Computer Programming in C",
        "title": "Lab Manual 02 — Structures, Unions & File Processing",
        "dept": "Department of Computer Engineering",
        "theory_questions": [
            "What is the difference between a structure and a union in C in terms of memory allocation?",
            "Explain structure padding and byte alignment in C compilers.",
            "Explain file opening modes ('r', 'w', 'a', 'r+', 'wb') in standard C library.",
            "How do fread() and fwrite() differ from fprintf() and fscanf()?"
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Student Record Management with Array of Structures",
                "text": "Write a C program to store records of N students (RollNo, Name, 3 Subject Marks) in an array of structures, calculate total and percentage for each student, and display a rank list."
            },
            {
                "num": 2,
                "title": "File Word Frequency and Character Counter",
                "text": "Write a C program to open a user-specified text file in read mode, count total characters, whitespace, lines, and words, and write summary results to a new output file."
            },
            {
                "num": 3,
                "title": "Binary File Record Storage using Structures",
                "text": "Write a C program to write employee records (ID, Name, Salary) to a binary file using fwrite() and read them back using fread() with search by employee ID."
            }
        ]
    },

    # ══════════════════════════════════════════════════════════════════════════
    # 4. C++ DSA FOCUS (3 Manuals)
    # ══════════════════════════════════════════════════════════════════════════
    {
        "folder": "c_cpp",
        "filename": "cpp_manual_1_stacks_and_queues.docx",
        "course": "CS204: Data Structures & Algorithms with C++",
        "title": "Lab Manual 01 — Linear Data Structures: Stacks & Queues",
        "dept": "Department of Computer Science & Engineering",
        "theory_questions": [
            "Explain the Last-In-First-Out (LIFO) and First-In-First-Out (FIFO) principles with real-world applications.",
            "What is stack overflow and stack underflow condition? How are they checked in array implementations?",
            "Why is a Circular Queue preferred over a Linear Queue implemented using a fixed-size array?",
            "Explain the algorithm to convert an Infix expression to Postfix notation using an operator stack."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Menu-Driven Stack with User Input",
                "text": "Write a C++ program to implement an interactive menu-driven Stack using a class with push, pop, peek, and display operations, taking continuous user inputs for elements until exit."
            },
            {
                "num": 2,
                "title": "Circular Queue Implementation with User Input",
                "text": "Write a C++ program to implement a Circular Queue of size N using arrays, providing an interactive console menu with enqueue, dequeue, front, and display operations."
            },
            {
                "num": 3,
                "title": "Infix to Postfix Converter and Evaluator",
                "text": "Write a C++ program that reads an infix arithmetic expression from standard input, converts it into postfix notation using a stack, and evaluates the postfix expression to compute the final value."
            }
        ]
    },
    {
        "folder": "c_cpp",
        "filename": "cpp_manual_2_trees_bst_traversals.docx",
        "course": "CS204: Data Structures & Algorithms with C++",
        "title": "Lab Manual 02 — Non-Linear Data Structures: Binary Search Trees",
        "dept": "Department of Computer Science & Engineering",
        "theory_questions": [
            "Define a Binary Search Tree (BST) and state its ordering property.",
            "Compare Inorder, Preorder, and Postorder traversals. Which traversal gives nodes in sorted order for a BST?",
            "Explain the three cases of node deletion in a Binary Search Tree (leaf node, one child, two children).",
            "What is tree height and balance factor? Explain how an unbalanced BST can degrade to O(N) search time."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Interactive Binary Search Tree with Traversals",
                "text": "Write a C++ program to construct a Binary Search Tree (BST) where nodes are dynamically inserted based on user input, and display Inorder, Preorder, Postorder, and Level-Order traversals."
            },
            {
                "num": 2,
                "title": "BST Search and Node Deletion with Interactive Menu",
                "text": "Write a C++ program implementing a menu-driven BST supporting user input to: (1) Insert Key, (2) Search Key, (3) Delete Key (handling 0, 1, or 2 children), and (4) Display Tree in Inorder."
            },
            {
                "num": 3,
                "title": "Tree Metrics: Height, Leaf Count, and Mirror Image",
                "text": "Write a C++ program to build a binary tree from user inputs and compute: (a) total node count, (b) total leaf nodes, (c) maximum tree height, and (d) convert the tree into its mirror image."
            }
        ]
    },
    {
        "folder": "c_cpp",
        "filename": "cpp_manual_3_graphs_dijkstra_traversals.docx",
        "course": "CS204: Data Structures & Algorithms with C++",
        "title": "Lab Manual 03 — Graph Algorithms: Traversals & Shortest Path",
        "dept": "Department of Computer Science & Engineering",
        "theory_questions": [
            "Compare adjacency matrix and adjacency list graph representations in terms of space and time complexity.",
            "Explain Breadth First Search (BFS) and Depth First Search (DFS) algorithms with their data structure requirements.",
            "What is Dijkstra's algorithm and what are its limitations regarding negative edge weights?",
            "Explain Cycle Detection in directed and undirected graphs."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Graph BFS and DFS Traversals with User Input",
                "text": "Write a C++ program to represent a graph using an adjacency list where the user enters the number of vertices and directed/undirected edges, and perform BFS starting from a user-selected source vertex and DFS traversals."
            },
            {
                "num": 2,
                "title": "Dijkstra's Single Source Shortest Path Algorithm",
                "text": "Write a C++ program to implement Dijkstra's algorithm using an adjacency matrix, reading the number of vertices, weighted edges from user input, and computing the minimum distance and shortest path from a user-specified source vertex to all other vertices."
            },
            {
                "num": 3,
                "title": "Prim's or Kruskal's Minimum Spanning Tree Algorithm",
                "text": "Write a C++ program to find the Minimum Spanning Tree (MST) of a connected weighted undirected graph entered by the user using Prim's algorithm, displaying the selected edges and total minimum weight."
            }
        ]
    },

    # ══════════════════════════════════════════════════════════════════════════
    # 5. WEB DEVELOPMENT (5 Manuals)
    # ══════════════════════════════════════════════════════════════════════════
    {
        "folder": "webdev",
        "filename": "webdev_manual_1_semantic_html_css_flex_grid.docx",
        "course": "CS305: Modern Web Development Laboratory",
        "title": "Lab Manual 01 — Semantic HTML5 & Modern CSS Layouts (Flexbox & Grid)",
        "dept": "Department of Software Engineering",
        "theory_questions": [
            "Explain semantic HTML5 tags (<header>, <nav>, <main>, <article>, <section>, <footer>) and their importance for SEO and accessibility.",
            "Differentiate between CSS Flexbox (1D) and CSS Grid (2D) layout models and when to use each.",
            "What is the CSS Box Model? Explain content, padding, border, margin, and box-sizing: border-box.",
            "Explain CSS media queries and mobile-first responsive web design principles."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Responsive Product Showcase with CSS Grid and Flexbox",
                "text": "Write an HTML5 and CSS3 webpage for an e-commerce product catalog using CSS Grid for the product cards and Flexbox for the navigation header, featuring responsive media queries for mobile, tablet, and desktop views."
            },
            {
                "num": 2,
                "title": "Accessible Student Registration Form with CSS Styling",
                "text": "Write an HTML5 form for college course registration using semantic tags, required input fields (name, email, date of birth, course select dropdown, radio buttons, file upload), styled cleanly with modern CSS."
            },
            {
                "num": 3,
                "title": "Sticky Responsive Navigation Bar with CSS Transitions",
                "text": "Write an HTML and CSS responsive sticky navigation bar that stays fixed at the top on scroll, features smooth hover transitions on navigation links, and highlights the active page link."
            }
        ]
    },
    {
        "folder": "webdev",
        "filename": "webdev_manual_2_javascript_dom_events_validation.docx",
        "course": "CS305: Modern Web Development Laboratory",
        "title": "Lab Manual 02 — JavaScript DOM Manipulation, Events & Form Validation",
        "dept": "Department of Software Engineering",
        "theory_questions": [
            "Explain the Document Object Model (DOM) tree and how JavaScript accesses DOM nodes.",
            "What is event bubbling, event capturing, and event delegation in modern JavaScript?",
            "Explain regular expressions in JavaScript and how test() and match() methods are used for validation.",
            "Differentiate between localStorage, sessionStorage, and cookies in client-side storage."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Real-Time Client-Side Form Validation with Regex",
                "text": "Write an HTML and JavaScript program for a user signup form that validates Name (min 3 chars), Email (standard regex), and Password (min 8 chars, 1 uppercase, 1 digit) in real-time as the user types, displaying inline error and success feedback."
            },
            {
                "num": 2,
                "title": "Interactive Calculator with Keyboard and Click Events",
                "text": "Write an HTML, CSS, and JavaScript program to build a functional arithmetic calculator with display screen, supporting both mouse clicks and keyboard event listeners for basic operations (+, -, *, /) and clear."
            },
            {
                "num": 3,
                "title": "Dynamic Filterable Table with LocalStorage",
                "text": "Write an HTML and JavaScript program to create an interactive student score table with search filter by name, sorting by score on header click, and saving new rows to localStorage."
            }
        ]
    },
    {
        "folder": "webdev",
        "filename": "webdev_manual_3_react_components_props_state.docx",
        "course": "CS305: Modern Web Development Laboratory",
        "title": "Lab Manual 03 — React Fundamentals: Components, Props & useState",
        "dept": "Department of Software Engineering",
        "theory_questions": [
            "What is JSX and how does React's Virtual DOM differ from the real browser DOM?",
            "Explain the difference between Props (immutable) and State (mutable) in React components.",
            "What are React Hooks? Explain the rules of hooks and the working of useState.",
            "Explain component re-rendering triggers and unidirectional data flow in React."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Interactive Counter App with Increment, Decrement, and Reset",
                "text": "Write a React component App.jsx using the useState hook to implement a counter with Increment, Decrement, Reset buttons, step size selection, and conditional color rendering when count is negative or zero."
            },
            {
                "num": 2,
                "title": "Dynamic Product Card Catalog with Props",
                "text": "Write a React application with a parent ProductList component passing product data (title, price, image, inStock status) via props to child ProductCard components, including an 'Add to Cart' badge handler."
            },
            {
                "num": 3,
                "title": "Tabbed Content Switcher Component",
                "text": "Write a React component that renders tab headers ('Overview', 'Specifications', 'Reviews') and conditionally renders the corresponding content section based on active tab state in useState."
            }
        ]
    },
    {
        "folder": "webdev",
        "filename": "webdev_manual_4_react_todo_localstorage.docx",
        "course": "CS305: Modern Web Development Laboratory",
        "title": "Lab Manual 04 — Advanced React: useEffect & LocalStorage Persistence",
        "dept": "Department of Software Engineering",
        "theory_questions": [
            "Explain the useEffect hook, its dependency array variants (none, [], [state]), and cleanup functions.",
            "What are controlled vs uncontrolled components in React form handling?",
            "How does React handle list rendering and why is a unique key prop mandatory for map()?",
            "Explain custom hooks in React and how they encourage logic reusability."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Persistent Todo List with Filter Tabs and LocalStorage",
                "text": "Write a React application TodoApp.jsx implementing full task management (add task, toggle complete, delete task, filter by All/Active/Completed) synchronized to browser localStorage using the useEffect hook."
            },
            {
                "num": 2,
                "title": "Live Character and Word Counter with Auto-Save",
                "text": "Write a React text editor component that updates character and word counts in real time, warns the user when exceeding 200 characters, and automatically saves draft text to localStorage using useEffect."
            },
            {
                "num": 3,
                "title": "Dark/Light Theme Toggle with Persistent Preference",
                "text": "Write a React application component with a Dark/Light mode theme switch that toggles CSS classes on the document body and saves the user's selected preference in localStorage."
            }
        ]
    },
    {
        "folder": "webdev",
        "filename": "webdev_manual_5_nodejs_express_rest_cookies.docx",
        "course": "CS305: Modern Web Development Laboratory",
        "title": "Lab Manual 05 — Backend Development: Node.js, Express REST API & Cookies",
        "dept": "Department of Software Engineering",
        "theory_questions": [
            "Explain the Node.js event loop, non-blocking I/O model, and single-threaded concurrency.",
            "What is REST architecture? Describe standard HTTP methods (GET, POST, PUT, DELETE) and status codes (200, 201, 400, 401, 404, 500).",
            "What is middleware in Express.js? Explain the role of req, res, and next() in the request-response lifecycle.",
            "Explain HTTP cookies security attributes: HttpOnly, Secure, SameSite, and maxAge."
        ],
        "lab_questions": [
            {
                "num": 1,
                "title": "Full CRUD Student REST API with Express",
                "text": "Write a Node.js and Express.js server (server.js) implementing RESTful CRUD endpoints for students (/api/students): GET all students, GET student by ID, POST new student with validation, PUT update student, and DELETE student with appropriate HTTP status codes."
            },
            {
                "num": 2,
                "title": "Session Authentication with Express-Session and Cookies",
                "text": "Write an Express.js backend (app.js) using express-session and cookie-parser to handle user login, issue a signed HttpOnly session cookie, protect a /profile endpoint to authenticated users only, and implement /logout to destroy the session."
            },
            {
                "num": 3,
                "title": "Static File Server and Request Logging Middleware",
                "text": "Write an Express.js application with custom logging middleware that logs HTTP method, URL, timestamp, and response status to console, serving static files from a public directory with express.static."
            }
        ]
    }
]


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def create_manual(data):
    doc = Document()

    # Set margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # 1. Header Box / Title
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    c.width = Inches(6.7)
    set_cell_background(c, "1E293B")
    set_cell_margins(c, top=160, bottom=160, left=200, right=200)

    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p.add_run("DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
    r_inst.font.name = "Segoe UI"
    r_inst.font.size = Pt(10)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    r_title = p.add_run(f"{data['title'].upper()}\n")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(14)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    r_sub = p.add_run(f"Course: {data['course']}  |  Academic Year: 2025–2026")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8)

    doc.add_paragraph()

    # 2. General Instructions / Objectives
    h_obj = doc.add_paragraph()
    r = h_obj.add_run("Course Objectives & Lab Guidelines")
    r.font.name = "Segoe UI"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.line_spacing = 1.15
    p_obj.paragraph_format.space_after = Pt(8)
    r_text = p_obj.add_run(
        "Students are required to solve all programming exercises individually during the lab session. "
        "Each program must include proper variable naming, boundary test case validation, and user documentation. "
        "Viva questions will be assessed based on theoretical understanding and practical demonstration."
    )
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # 3. SECTION A: Theory & Viva Questions (EXCLUDED from code generation by design)
    h_secA = doc.add_paragraph()
    r = h_secA.add_run("Section A: Theory & Viva Voce Questions")
    r.font.name = "Segoe UI"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)  # Red tone to mark theory
    p_note = doc.add_paragraph()
    r_note = p_note.add_run("(These conceptual questions test theoretical foundations and should be answered in the lab record.)")
    r_note.font.name = "Segoe UI"
    r_note.font.size = Pt(8.5)
    r_note.font.italic = True
    r_note.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    for i, t_q in enumerate(data["theory_questions"], 1):
        p_t = doc.add_paragraph()
        p_t.paragraph_format.left_indent = Inches(0.2)
        p_t.paragraph_format.space_after = Pt(4)
        r_num = p_t.add_run(f"Q{i}. ")
        r_num.font.name = "Segoe UI"
        r_num.font.size = Pt(9.5)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        r_body = p_t.add_run(t_q)
        r_body.font.name = "Segoe UI"
        r_body.font.size = Pt(9.5)
        r_body.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph()

    # 4. SECTION B: Laboratory Programming Exercises (ACTIONABLE questions to be extracted & executed)
    h_secB = doc.add_paragraph()
    r = h_secB.add_run("Section B: Laboratory Programming Exercises")
    r.font.name = "Segoe UI"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x16, 0x65, 0x34)  # Green tone for actionable lab tasks

    p_note_b = doc.add_paragraph()
    r_note_b = p_note_b.add_run("(Implement, compile, execute and capture program outputs for each of the following exercises.)")
    r_note_b.font.name = "Segoe UI"
    r_note_b.font.size = Pt(8.5)
    r_note_b.font.italic = True
    r_note_b.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    for item in data["lab_questions"]:
        # Container box for each experiment
        t_exp = doc.add_table(rows=1, cols=1)
        t_exp.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t_exp.cell(0, 0)
        cell.width = Inches(6.7)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

        cp = cell.paragraphs[0]
        r_num = cp.add_run(f"Experiment {item['num']}: {item['title']}\n")
        r_num.font.name = "Segoe UI"
        r_num.font.size = Pt(10.5)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        r_aim = cp.add_run("Aim / Problem Statement: ")
        r_aim.font.name = "Segoe UI"
        r_aim.font.size = Pt(9.5)
        r_aim.font.bold = True
        r_aim.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

        r_desc = cp.add_run(f"{item['text']}\n")
        r_desc.font.name = "Segoe UI"
        r_desc.font.size = Pt(9.5)
        r_desc.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        r_sub = cp.add_run("Submission Requirements: ")
        r_sub.font.name = "Segoe UI"
        r_sub.font.size = Pt(8.5)
        r_sub.font.bold = True
        r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        r_req = cp.add_run("Source code implementation, compiler/interpreter execution screenshot, and verified test cases.")
        r_req.font.name = "Segoe UI"
        r_req.font.size = Pt(8.5)
        r_req.font.italic = True
        r_req.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

        doc.add_paragraph()

    # Save
    out_dir = os.path.join(OUTPUT_BASE, data["folder"])
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, data["filename"])
    doc.save(out_path)
    size_kb = os.path.getsize(out_path) // 1024
    print(f"  [OK] Saved: {out_path} ({size_kb} KB)")
    return out_path


def main():
    print("=" * 65)
    print("GENERATING 20 FINAL BOSS LEVEL LAB MANUALS (.DOCX)")
    print("=" * 65)
    created = []
    for item in MANUALS_DATA:
        p = create_manual(item)
        created.append(p)
    print("=" * 65)
    print(f"SUCCESS: Generated {len(created)} manuals across 4 directories.")
    print("=" * 65)


if __name__ == "__main__":
    main()
