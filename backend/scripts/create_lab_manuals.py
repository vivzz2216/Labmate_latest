import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def style_heading(p, text, size=16, bold=True, color=(30, 58, 138)):
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)

def add_paragraph(doc, text, bold_prefix="", size=11, italic=False):
    p = doc.add_paragraph()
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(size)
        r_pre.font.bold = True
    r_body = p.add_run(text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(size)
    r_body.font.italic = italic
    return p

def create_manual(filepath, title, course_code, theory_questions, lab_exercises):
    doc = Document()
    
    # Document header
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING\n")
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"{title.upper()}\nCourse Code: {course_code}")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(12)
    r_title.font.bold = True
    
    doc.add_paragraph() # Spacing

    # Section 1: Objectives & Outcomes (Ignored by parser)
    p_obj = doc.add_paragraph()
    style_heading(p_obj, "Course Objectives & Outcomes", size=13, color=(51, 65, 85))
    add_paragraph(doc, "Understand fundamental paradigms, syntax, design patterns, and application engineering.", bold_prefix="Objectives: ")
    add_paragraph(doc, "Students will be capable of designing, debugging, and compiling software solutions.", bold_prefix="Learning Outcomes: ")
    
    doc.add_paragraph()

    # Section 2: Theory Questions (Must be IGNORED by LabMate parser)
    p_theory = doc.add_paragraph()
    style_heading(p_theory, "Theory & Viva Questions", size=13, color=(185, 28, 28))
    for idx, tq in enumerate(theory_questions, 1):
        add_paragraph(doc, f"{tq['q']}\nAnswer: {tq['a']}", bold_prefix=f"Q{idx}. ")

    doc.add_paragraph()

    # Section 3: Laboratory Exercises (Must be EXTRACTED by LabMate parser)
    p_lab = doc.add_paragraph()
    style_heading(p_lab, "Laboratory Exercises & Programming Tasks", size=13, color=(21, 128, 61))
    for idx, prog in enumerate(lab_exercises, 1):
        add_paragraph(doc, prog['description'], bold_prefix=f"Program {idx}: ")
        if 'aim' in prog:
            add_paragraph(doc, prog['aim'], bold_prefix="Aim: ", italic=True)

    # Ensure output directory exists
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    doc.save(filepath)
    print(f"Created: {filepath}")

def main():
    base_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "sample_lab_manuals")
    os.makedirs(base_dir, exist_ok=True)

    # 1. Python Lab Manual
    create_manual(
        filepath=os.path.join(base_dir, "python_lab_manual.docx"),
        title="Python Programming Laboratory",
        course_code="CS301-PY",
        theory_questions=[
            {"q": "Explain the Global Interpreter Lock (GIL) and its implications on concurrency in Python.", "a": "GIL is a mutex that protects access to Python objects, preventing multiple threads from executing Python bytecodes simultaneously."},
            {"q": "Differentiate between mutable and immutable data types in Python with examples.", "a": "Mutable types (lists, dicts, sets) can be altered in place. Immutable types (tuples, strings, ints) cannot be changed after creation."},
            {"q": "What are Python decorators and how do they function?", "a": "Decorators are functions that take another function as an argument and extend its behavior without explicitly modifying it."},
            {"q": "Describe the difference between shallow copy and deep copy.", "a": "A shallow copy constructs a new compound object and inserts references. A deep copy recursively copies all objects found."}
        ],
        lab_exercises=[
            {
                "description": "Write a Python program to calculate the factorial of an integer provided by the user using recursion, with validation for negative numbers.",
                "aim": "Implement a recursive factorial function and demonstrate error checking for edge conditions."
            },
            {
                "description": "Write a Python program to read numbers from a text file, calculate their sum, average, and standard deviation, and write the formatted results to an output file.",
                "aim": "Demonstrate robust file handling using context managers and statistical computations."
            }
        ]
    )

    # 2. Java Lab Manual
    create_manual(
        filepath=os.path.join(base_dir, "java_lab_manual.docx"),
        title="Object Oriented Programming with Java Laboratory",
        course_code="CS302-JAVA",
        theory_questions=[
            {"q": "What is the difference between JVM, JRE, and JDK?", "a": "JDK is the development kit including compiler and tools; JRE is the runtime environment; JVM is the virtual machine executing bytecode."},
            {"q": "Explain method overloading vs method overriding in Java.", "a": "Overloading occurs at compile-time with different signatures in the same class. Overriding occurs at runtime in a subclass with the same signature."},
            {"q": "Why is the String class immutable in Java?", "a": "For security, synchronization, caching in the String pool, and performance optimization."},
            {"q": "Explain the life cycle of a Thread in Java.", "a": "New, Runnable, Blocked, Waiting, Timed Waiting, and Terminated."}
        ],
        lab_exercises=[
            {
                "description": "Write a Java program to accept student details (name, roll number, and marks in three subjects) using Scanner, compute total and percentage, and display grade based on criteria.",
                "aim": "Implement interactive console I/O using java.util.Scanner and conditional branching."
            },
            {
                "description": "Write a Java program to demonstrate multi-level inheritance by creating a Person base class, Employee subclass, and Manager derived class with salary computation.",
                "aim": "Implement inheritance hierarchy, constructor chaining with super(), and method overriding."
            }
        ]
    )

    # 3. C Lab Manual
    create_manual(
        filepath=os.path.join(base_dir, "c_lab_manual.docx"),
        title="Problem Solving with C Programming Laboratory",
        course_code="CS303-C",
        theory_questions=[
            {"q": "What is the difference between malloc() and calloc() in C?", "a": "malloc() allocates uninitialized memory; calloc() allocates memory and initializes all bytes to zero."},
            {"q": "Explain pointer arithmetic and how array indexing relates to pointers.", "a": "Array indexing a[i] is internally evaluated as *(a + i) using base address and type sizing."},
            {"q": "What are storage classes in C? Explain auto, static, extern, and register.", "a": "Storage classes define the scope, visibility, and lifetime of variables and functions."},
            {"q": "Explain the difference between passing by value and passing by reference using pointers.", "a": "Pass by value sends a copy; pass by reference sends memory address allowing in-place modification."}
        ],
        lab_exercises=[
            {
                "description": "Write a C program to implement the Bubble Sort algorithm on an array of integers and display the sorted array.",
                "aim": "Demonstrate array manipulation, nested loops, and swapping logic in C."
            },
            {
                "description": "Write a C program using pointers to check whether a given string is a palindrome.",
                "aim": "Implement pointer arithmetic to traverse and compare strings without string library functions."
            }
        ]
    )

    # 4. C++ Lab Manual
    create_manual(
        filepath=os.path.join(base_dir, "cpp_lab_manual.docx"),
        title="Object Oriented Programming in C++ Laboratory",
        course_code="CS304-CPP",
        theory_questions=[
            {"q": "What is the role of the virtual destructor in C++?", "a": "A virtual destructor ensures that the derived class destructor is called when deleting an object through a base class pointer."},
            {"q": "Explain the difference between struct and class in C++.", "a": "In struct, default member and inheritance access is public; in class, it is private."},
            {"q": "What are friend functions and friend classes in C++?", "a": "A friend function or class has access to the private and protected members of the declaring class."},
            {"q": "Explain copy constructor and when it is implicitly invoked.", "a": "It initializes an object using an existing object of the same class, called during pass-by-value or explicit copy initialization."}
        ],
        lab_exercises=[
            {
                "description": "Write a C++ program to create a Matrix class to add and multiply two 2x2 matrices using member functions.",
                "aim": "Demonstrate class declaration, private data members, and object operations in C++."
            },
            {
                "description": "Write a C++ program to demonstrate operator overloading by overloading the '+' operator to add two Complex numbers.",
                "aim": "Implement operator+ function and member function overloading for user-defined types."
            }
        ]
    )

    # 5. Web Development Lab Manual
    create_manual(
        filepath=os.path.join(base_dir, "webdev_lab_manual.docx"),
        title="Web Technologies & Frontend Development Laboratory",
        course_code="CS305-WEB",
        theory_questions=[
            {"q": "Explain the CSS Box Model and the difference between content-box and border-box.", "a": "The box model includes content, padding, border, and margin. border-box includes padding and border in total width/height."},
            {"q": "What is the difference between localStorage, sessionStorage, and cookies?", "a": "localStorage persists indefinitely, sessionStorage lasts for the session, and cookies are sent with every HTTP request with expiration."},
            {"q": "Explain event bubbling and event capturing in JavaScript.", "a": "In bubbling, events propagate from inner target outward; in capturing, events travel downward from window to target."},
            {"q": "Describe the Virtual DOM in React and how reconciliation works.", "a": "Virtual DOM is an in-memory representation. Reconciliation diffs virtual trees and applies minimal DOM updates."}
        ],
        lab_exercises=[
            {
                "description": "Write an HTML and CSS program to create a modern responsive product card with hover animations and responsive typography.",
                "aim": "Implement semantic HTML5 tags and modern CSS Flexbox styling."
            },
            {
                "description": "Write an interactive JavaScript application to create a dynamic counter widget with increment, decrement, and reset functionality.",
                "aim": "Demonstrate DOM querySelector, addEventListener, and dynamic state update."
            }
        ]
    )

if __name__ == "__main__":
    main()
