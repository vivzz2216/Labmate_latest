"""Replace only reviewed evidence images in existing Word reports.

Run this with the bundled Codex document Python after
repair_assignment_evidence.py has produced a complete manifest. Original
downloads and uploaded manuals are never overwritten.
"""

import argparse
import io
import json
import re
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph as WordParagraph
from PIL import Image


QUESTION = re.compile(r"^Question (\d+):\s")


def corrected_experiment_two_page(destination: Path):
    """Re-typeset only the erroneous source PDF page for the Word copy."""
    pdf_path = destination / "experiment_2_corrected_source_page.pdf"
    image_path = destination / "experiment_2_corrected_source_page.png"
    if image_path.is_file():
        return image_path
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    styles = getSampleStyleSheet()
    heading = ParagraphStyle("ManualHeading", parent=styles["Heading2"], textColor=colors.black,
                             fontName="Helvetica-Bold", fontSize=12, leading=15, spaceBefore=11, spaceAfter=5)
    title = ParagraphStyle("ManualTitle", parent=styles["Title"], textColor=colors.black,
                           fontName="Helvetica-Bold", fontSize=17, leading=21, alignment=TA_CENTER, spaceAfter=9)
    body = ParagraphStyle("ManualBody", parent=styles["Normal"], fontName="Helvetica",
                          fontSize=9.5, leading=13, spaceAfter=4)
    small = ParagraphStyle("ManualSmall", parent=body, fontSize=9, leading=12)
    story = [
        Paragraph("St Francis Institute of Technology", title),
        Paragraph("Department of Information Technology | Python Lab | Experiment 2", body),
        Paragraph("Python Inbuilt Mathematical and String Methods", heading),
        Paragraph("Aim: Write Python programs demonstrating ten math functions and ten string methods.", body),
        Paragraph("Objectives: Understand the use and return values of mathematical functions and string methods.", body),
        Paragraph("Prerequisite: Python basics. Requirements: PC, Python 3, and IDLE or another Python IDE.", body),
        Paragraph("Pre Experiment Theory", heading),
        Paragraph("Math methods: Python provides mathematical functions for numeric calculations, including trigonometric functions. The random module provides pseudorandom values useful in simulations and testing.", body),
        Paragraph("Strings in Python: Strings are immutable, so string methods do not modify the original string. Their return types vary: upper() returns a new string, find() returns an integer index, and startswith() returns a Boolean.", body),
        Paragraph("Operators in Python", heading),
        Paragraph("Arithmetic: +, -, *, /, //, %, **. Assignment: =, +=, -=, *=, /=, //=, %=, **=, &=, |=, ^=, <<=, >>=.", small),
        Paragraph("Comparison: ==, !=, &lt;, &gt;, &lt;=, &gt;=. Logical: and, or, not. Identity: is, is not. Membership: in, not in. Bitwise: &amp;, |, ^, ~, &lt;&lt;, &gt;&gt;.", small),
        Paragraph("Laboratory Exercise", heading),
        Paragraph("Open a Python editor, create and save a .py file, run the program, inspect actual output, and test multiple cases. Add relevant comments to the program.", body),
        Paragraph("Post Experiment Exercise", heading),
        Paragraph("List and explain the building blocks of a Python program. The remaining original manual page follows unchanged in this report.", body),
    ]
    SimpleDocTemplate(str(pdf_path), pagesize=A4, leftMargin=52, rightMargin=52,
                      topMargin=48, bottomMargin=48).build(story)
    document = pdfium.PdfDocument(str(pdf_path))
    if len(document) != 1:
        raise ValueError("Corrected theory page did not fit on one page.")
    document[0].render(scale=2).to_pil().save(image_path)
    document.close()
    return image_path


def image_paragraph(document, after, image_path, maximum_width=Inches(6.2)):
    """Insert a readable screenshot, splitting tall captures losslessly."""
    section = document.sections[-1]
    usable_width = section.page_width - section.left_margin - section.right_margin
    usable_height = section.page_height - section.top_margin - section.bottom_margin - Inches(0.65)
    display_width = min(usable_width, maximum_width)
    with Image.open(image_path) as image:
        width, height = image.size
        max_slice_height = max(1, int(width * usable_height / display_width))

        # Proportional scaling to avoid unnecessary slices for images within 25% of page height
        if max_slice_height < height <= int(max_slice_height * 1.25):
            scaled_w = int(width * usable_height / height)
            if scaled_w >= Inches(4.5):
                display_width = scaled_w
                max_slice_height = height

        slices = []
        top = 0
        while top < height:
            bottom = min(height, top + max_slice_height)
            if bottom < height:
                for candidate in range(bottom, max(top + 1, bottom - 150), -1):
                    sample = image.crop((min(40, width - 1), candidate, max(min(41, width), width - 50), candidate + 1)).convert("L")
                    if sample.getextrema()[0] >= 240:
                        bottom = candidate
                        break
            slices.append((top, bottom))
            top = bottom

        total_slices = len(slices)
        for slice_idx, (s_top, s_bottom) in enumerate(slices):
            element = OxmlElement("w:p")
            after.addnext(element)
            paragraph = WordParagraph(element, document)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_together = True
            if slice_idx > 0:
                paragraph.paragraph_format.page_break_before = True

            if s_top == 0 and s_bottom == height:
                source = str(image_path)
            else:
                source = io.BytesIO()
                image.crop((0, s_top, width, s_bottom)).save(source, format="PNG")
                source.seek(0)

            slice_h = int(display_width * (s_bottom - s_top) / width)
            paragraph.add_run().add_picture(source, width=int(display_width), height=slice_h)
            after = element

            if total_slices > 1:
                cap_el = OxmlElement("w:p")
                after.addnext(cap_el)
                cap_para = WordParagraph(cap_el, document)
                cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_para.paragraph_format.keep_together = True
                cap_text = f"Code (Part {slice_idx + 1} of {total_slices})" if slice_idx == 0 else f"Code (Part {slice_idx + 1} of {total_slices} — continued)"
                run = cap_para.add_run(cap_text)
                run.font.size = Pt(9)
                run.italic = True
                after = cap_el

    return after


def replace_question_images(document, question_id, screenshot_paths):
    paragraphs = list(document.paragraphs)
    matching = [(index, paragraph) for index, paragraph in enumerate(paragraphs)
                if QUESTION.match(paragraph.text) and int(QUESTION.match(paragraph.text).group(1)) == question_id]
    if len(matching) != 1:
        raise ValueError(f"Expected one generated Question {question_id}; found {len(matching)}")
    index, question = matching[0]
    question.paragraph_format.keep_together = True
    question.paragraph_format.keep_with_next = True
    following = next((next_index for next_index in range(index + 1, len(paragraphs))
                      if QUESTION.match(paragraphs[next_index].text)), len(paragraphs))
    old_images = [paragraph for paragraph in paragraphs[index + 1:following]
                  if paragraph._element.xpath(".//a:blip")]
    if not old_images:
        raise ValueError(f"No prior screenshots found for Question {question_id}.")
    for paragraph in old_images:
        paragraph._element.getparent().remove(paragraph._element)
    anchor = question._element
    for screenshot in screenshot_paths:
        path = Path(screenshot)
        if not path.is_file():
            raise ValueError(f"Verified screenshot missing: {path}")
        anchor = image_paragraph(document, anchor, path)
    return len(old_images)


def replace_first_pdf_page(document, corrected_image):
    first = next((paragraph for paragraph in document.paragraphs
                  if paragraph._element.xpath(".//a:blip")), None)
    if first is None:
        raise ValueError("Original PDF page image was not found.")
    first._element.getparent().remove(first._element)
    anchor = document.paragraphs[1]._element
    image_paragraph(document, anchor, corrected_image, maximum_width=Inches(5.7))


def main(output_dir, original_dir):
    manifest = json.loads((output_dir / "manifest.json").read_text(encoding="utf-8"))
    corrected_theory = corrected_experiment_two_page(output_dir)
    completed = []
    for report in manifest["reports"]:
        name = report["name"]
        source = original_dir / name
        if not source.is_file():
            raise FileNotFoundError(source)
        document = Document(source)
        if report["workflow_id"] == 9:
            replace_first_pdf_page(document, corrected_theory)
        for result in report["results"]:
            evidence = result.get("screenshot_paths", [])
            if not evidence or not all(str(path).startswith(str(output_dir / "evidence")) for path in evidence):
                continue
            replace_question_images(document, result["id"], evidence)
            if report["workflow_id"] == 8 and result["id"] == 3:
                question = next(paragraph for paragraph in document.paragraphs
                                if paragraph.text.startswith("Question 3:"))
                question.text = ("Question 3: The manual requests ten tuple methods, but Python tuples have "
                                 "only count() and index(). The screenshots demonstrate both methods and "
                                 "eight related operations or functions, clearly distinguished.")
        target = output_dir / name
        document.save(target)
        reopened = Document(target)
        if not reopened.inline_shapes:
            raise ValueError(f"No images remain in {target}")
        completed.append({"file": str(target), "images": len(reopened.inline_shapes)})
        print(f"Built {target.name} with {len(reopened.inline_shapes)} visible images", flush=True)
    (output_dir / "report_inventory.json").write_text(json.dumps(completed, indent=2), encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--original-dir", type=Path, required=True)
    args = parser.parse_args()
    main(args.output_dir.resolve(), args.original_dir.resolve())
