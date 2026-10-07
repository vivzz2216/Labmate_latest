import docx
import os
import sys

# Path to the file
path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "Python_Functions_Laboratory_Manual.docx"))
print(f"Checking file: {path} (exists: {os.path.exists(path)})")

doc = docx.Document(path)
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n--- ALL PARAGRAPHS WITH CONTENT ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt:
        print(f"P{i:03d} (style={p.style.name if p.style else 'None'}): {txt}")

if doc.tables:
    print("\n--- TABLES ---")
    for t_idx, table in enumerate(doc.tables):
        print(f"\nTable {t_idx} ({len(table.rows)} rows):")
        for r_idx, row in enumerate(table.rows):
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            print(f"  R{r_idx}: {cells}")
